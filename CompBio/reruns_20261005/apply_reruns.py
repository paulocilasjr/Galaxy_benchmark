"""Swap the 53 reviewed CompBioBench Galaxy reruns into the task packages and the run-link workbook.

The reruns are first collected with the repository's own collector (analysis_execution/collect.py) into a staging
directory laid out like a task package. For every rerun this script then:
  1. moves the original run's trace folder, trace manifests, Galaxy history snapshot, selected outputs, job ledger and
     recovered code to analysis/<task>/versions/pre_rerun_20261005/ (nothing is deleted), after copying the task-level
     derived files there once;
  2. copies the staged rerun files into the same places;
  3. points the run's workbook row at the rerun trace and history (the original workbook is kept beside this script);
  4. rebuilds the affected task packages offline with analysis_execution/run.py and scripts/enrich_compbio_evidence.py.
An offline rebuild of an unchanged package reproduces its evidence exactly (only the audit timestamp differs).

Usage: python apply_reruns.py STAGE_DIR [--dry-run]
"""
import argparse
import csv
import json
import shutil
from pathlib import Path

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
ANALYSIS = BASE / "analysis"
WORKBOOK = BASE / "compbiobench_execution_condition_links.xlsx"
WORKBOOK_ORIGINAL = HERE / "compbiobench_execution_condition_links.pre_rerun_20261005.xlsx"
ARCHIVE = "versions/pre_rerun_20261005"
MODEL_LABEL = {"GPT-5.5": "Codex GPT-5.5", "GPT-5.6 Sol": "Codex GPT-5.6 Sol", "GPT-5.6 Luna": "Codex GPT-5.6 Luna",
               "DeepSeek V4 Pro": "Codex + DeepSeek-v4-pro-0813"}
TASK_FILES = ["README.md", "history_analysis.md", "history_analysis_evidence.json", "run_manifest.json",
              "input_manifest.json", "source_snapshots/huggingface_traces/manifest_galaxy.json",
              "source_snapshots/galaxy/retrieval_manifest.json", "recovered_code/manifest.json",
              "selected_outputs/galaxy/manifest.json"]


def galaxy_url(history_id):
    return f"https://usegalaxy.org/published/history?id={history_id}"


def move(src, dst, dry):
    if not src.exists():
        return
    print(f"  move {src.relative_to(ANALYSIS)} -> {dst.relative_to(ANALYSIS)}")
    if not dry:
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(src), str(dst))


def copy(src, dst, dry):
    print(f"  copy staged {dst.relative_to(ANALYSIS)}")
    if not dry:
        dst.parent.mkdir(parents=True, exist_ok=True)
        (shutil.copytree if src.is_dir() else shutil.copy2)(src, dst)


def swap_run(row, stage, dry):
    task, rid, new_hid = row["task"], row["run_id"], row["history_id_rerun"]
    pkg, staged = ANALYSIS / task, stage / task
    archive = pkg / ARCHIVE
    record = archive / "swapped_runs.json"
    done = json.loads(record.read_text()) if record.exists() else {}
    if rid in done:
        print(f"{task} {rid}: already swapped")
        return
    for name in TASK_FILES + [f"{task}.json"]:
        if (pkg / name).exists() and not (archive / name).exists():
            print(f"  keep task file {name}")
            if not dry:
                (archive / name).parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(pkg / name, archive / name)
    evidence = json.loads((pkg / "history_analysis_evidence.json").read_text())
    run = next(r for r in evidence["runs"] if r["run_id"] == rid)
    old_sources = [s for s in run["source_ids"] if s.startswith("src_galaxy_")]
    old_hid = old_sources[0][len("src_galaxy_"):] if old_sources else None
    if old_hid and any(f"src_galaxy_{old_hid}" in r["source_ids"] for r in evidence["runs"] if r["run_id"] != rid):
        raise RuntimeError(f"{task}: history {old_hid} is shared with another run; refusing to archive it")
    trace_source = next(s for s in evidence["sources"] if s["source_id"] == f"src_trace_{rid}")
    print(f"{task} {rid}: history {old_hid} -> {new_hid}")
    hf = "source_snapshots/huggingface_traces"
    for rel in (f"{hf}/files/{rid}", f"{hf}/manifests/{rid}.json", f"{hf}/indexes/{rid}.json",
                f"job_ledgers/galaxy/{rid}.json", f"recovered_code/galaxy/{rid}"):
        move(pkg / rel, archive / rel, dry)
    if old_hid:
        for rel in (f"source_snapshots/galaxy/{old_hid}", f"selected_outputs/galaxy/{old_hid}"):
            move(pkg / rel, archive / rel, dry)
    for rel in (f"{hf}/files/{rid}", f"{hf}/manifests/{rid}.json", f"{hf}/indexes/{rid}.json",
                f"source_snapshots/galaxy/{new_hid}", f"selected_outputs/galaxy/{new_hid}"):
        if (staged / rel).exists():
            copy(staged / rel, pkg / rel, dry)
        elif "selected_outputs" not in rel:
            raise FileNotFoundError(staged / rel)
    if not dry:
        done[rid] = {"original_trace_url": trace_source["location"], "original_history_id": old_hid,
                     "rerun_trace_url": row["trace_url_rerun"], "rerun_history_id": new_hid,
                     "reason": row["reason"]}
        record.write_text(json.dumps(done, indent=2, sort_keys=True) + "\n")


def update_workbook(rows, dry):
    import openpyxl
    if not WORKBOOK_ORIGINAL.exists() and not dry:
        shutil.copy2(WORKBOOK, WORKBOOK_ORIGINAL)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["CompBioBench links"]
    header = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
    col = {k: header[k] for k in ("task", "model", "condition", "replicate", "run_trace", "galaxy_history")}
    wanted = {(r["task"], MODEL_LABEL[r["cfg"]], int(r["replicate"])): r for r in rows}
    hits = 0
    for line in ws.iter_rows(min_row=2):
        cell = {k: line[c - 1] for k, c in col.items()}
        key = (str(cell["task"].value).strip(), str(cell["model"].value).strip(),
               int(str(cell["replicate"].value).strip().split()[-1].split("_")[-1]))
        if key not in wanted or "galaxy" not in str(cell["condition"].value).lower():
            continue
        r = wanted.pop(key)
        hits += 1
        for name, url in (("run_trace", r["trace_url_rerun"]), ("galaxy_history", galaxy_url(r["history_id_rerun"]))):
            target = cell[name]
            if target.hyperlink is not None:
                target.hyperlink = url
            target.value = url
    if wanted:
        raise RuntimeError(f"Workbook rows not found: {sorted(wanted)}")
    print(f"workbook: {hits} Galaxy rows point to reruns")
    if not dry:
        wb.save(WORKBOOK)


def rebuild(tasks):
    """Rebuild the affected task packages offline, then re-apply the CompBio enrichment and the overview link."""
    import subprocess
    import sys
    repo = BASE.parent
    cmd = [sys.executable, "run.py", str(WORKBOOK), "--benchmark", "compbio", "--offline"]
    for task in tasks:
        cmd += ["--task", task]
    subprocess.run(cmd, cwd=repo / "analysis_execution", check=True)
    sys.path.insert(0, str(repo / "scripts"))
    import enrich_compbio_evidence
    for task in tasks:
        enrich_compbio_evidence.enrich(ANALYSIS / task)
        report = ANALYSIS / task / "history_analysis.md"
        report.write_text(report.read_text().replace("`CompBio/compBio_overview.md`",
                                                     "[the benchmark overview](../../compBio_overview.md)"))
    print(f"rebuilt {len(tasks)} task packages")


def main():
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("stage", type=Path)
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()
    rows = list(csv.DictReader((HERE / "rerun_manifest.csv").open()))
    assert len(rows) == 53
    for row in rows:
        swap_run(row, a.stage.resolve(), a.dry_run)
    update_workbook(rows, a.dry_run)
    tasks = sorted({r["task"] for r in rows})
    print("tasks:", " ".join(tasks))
    if not a.dry_run:
        rebuild(tasks)


if __name__ == "__main__":
    main()
