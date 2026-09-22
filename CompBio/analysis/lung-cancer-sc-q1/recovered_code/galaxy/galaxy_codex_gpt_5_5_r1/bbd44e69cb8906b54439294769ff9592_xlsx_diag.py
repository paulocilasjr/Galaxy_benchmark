import sys, json
from openpyxl import load_workbook
path,out=sys.argv[1],sys.argv[2]
wb=load_workbook(path, read_only=True, data_only=True)
with open(out,'w') as fh:
    fh.write('sheets\t%s\n' % json.dumps(wb.get_sheet_names()))
    for ws in wb.worksheets:
        fh.write('SHEET\t%s\t%d\t%d\n' % (ws.title, ws.max_row, ws.max_column))
        rows=[]
        for i,row in enumerate(ws.iter_rows(), start=1):
            vals=[c.value for c in row]
            rows.append(vals)
            if i>=40: break
        fh.write('ROWS\t%s\t%s\n' % (ws.title, json.dumps(rows, ensure_ascii=True, default=str)))
