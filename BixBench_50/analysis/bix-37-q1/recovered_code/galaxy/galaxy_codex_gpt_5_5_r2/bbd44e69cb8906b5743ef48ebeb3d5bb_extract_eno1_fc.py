from __future__ import print_function
import sys
from openpyxl import load_workbook

def values(row):
    return [cell.value for cell in row]

infile = sys.argv[1]
outfile = 'eno1_fc.tsv'
wb = load_workbook(infile, read_only=True, data_only=True)
sheet_name = 'Tumor vs Normal'
if sheet_name not in wb.get_sheet_names():
    raise SystemExit('Required sheet not found: ' + sheet_name)
ws = wb.get_sheet_by_name(sheet_name)
rows = ws.iter_rows()
try:
    headers = [str(x).strip() if x is not None else '' for x in values(next(rows))]
except StopIteration:
    raise SystemExit('Workbook sheet is empty')
col = {h: i for i, h in enumerate(headers)}
required = ['protein', 'Description', 'gene', 'gene_id', 'Normal', 'Tumor', 'Ratio', 'FC', 'log2FC', 'compare']
missing = [h for h in required if h not in col]
if missing:
    raise SystemExit('Missing required columns: ' + ','.join(missing))
matched = []
for raw_row in rows:
    row = values(raw_row)
    gene = row[col['gene']]
    compare = row[col['compare']]
    if gene == 'ENO1' and compare == 'Tumor vs Normal':
        matched.append(row)
if len(matched) != 1:
    raise SystemExit('Expected one ENO1 Tumor vs Normal row, found {0}'.format(len(matched)))
row = matched[0]
with open(outfile, 'w') as out:
    out.write('gene\tprotein\tNormal\tTumor\tRatio\tFC\tlog2FC\n')
    out.write('{0}\t{1}\t{2}\t{3}\t{4}\t{5}\t{6}\n'.format(
        row[col['gene']], row[col['protein']], row[col['Normal']], row[col['Tumor']],
        row[col['Ratio']], row[col['FC']], row[col['log2FC']]))
