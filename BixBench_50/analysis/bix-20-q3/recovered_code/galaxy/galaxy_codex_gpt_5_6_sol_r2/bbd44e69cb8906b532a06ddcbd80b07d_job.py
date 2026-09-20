from __future__ import print_function
import csv
import re
import sys
from collections import Counter
from io import BytesIO
from openpyxl import load_workbook

SAMPLE_IDS = ["179","184","185","285","286","287","353","354","360","364","380","381","396","397","409","488","489","490","498","499","500","502","503","504","556","557","613","614","615","SRR5456220","SRR5456232","SRR5462651","SRR5462820","SRR5462858","SRR5462872","SRR5462955","SRR5462967","SRR5463095","SRR5463149","SRR5463270","SRR5463506","SRR5466052","SRR5466100","SRR5469760","SRR5469979","SRR5470425","SRR5550148","SRR5550381","SRR5550774","SRR5550794","SRR5551062","SRR5553488","SRR5555687","SRR5556325","SRR5560909","SRR5560927","SRR5561089","SRR5561107","SRR5561237","SRR5561287","SRR5561319","SRR5561584","SRR5561706","SRR5561875","SRR5561934","SRR5561950","SRR5562000","SRR5562034","SRR5562184","SRR5562220","SRR5562238","SRR5565258","SRR5565278","SRR5565931","SRR5565968","SRR5566591","SRR5566906","SRR5567694","SRR5567876","SRR5567954","SRR5571393","SRR5688631","SRR5688669","SRR5688704","SRR5688812","SRR6447679"]

def norm_id(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def truthy(value):
    return str(value).strip().lower() in ("true", "1", "yes")

def values(row):
    return [c.value for c in row]

def first_sheet(path):
    with open(path, "rb") as handle:
        wb = load_workbook(BytesIO(handle.read()), read_only=False, data_only=True)
    return wb[wb.sheetnames[0]]

def index_of(headers, name):
    for i, h in enumerate(headers):
        if str(h).strip() == name:
            return i
    raise RuntimeError("Missing required column: " + name)

def tokens(text):
    return [x.upper() for x in re.findall(r"[A-Za-z0-9_-]+", str(text or ""))]

status_path = sys.argv[1]
genes_path = sys.argv[2]
variant_paths = sys.argv[3:]
if len(variant_paths) != len(SAMPLE_IDS):
    raise RuntimeError("Expected %d variant inputs, received %d" % (len(SAMPLE_IDS), len(variant_paths)))

status_ws = first_sheet(status_path)
status_headers = values(next(status_ws.iter_rows(min_row=1, max_row=1)))
sid_i = index_of(status_headers, "Sample ID")
blm_i = index_of(status_headers, "BLM Mutation Status")
carrier_ids = set()
for row in status_ws.iter_rows(min_row=2):
    vals = values(row)
    if len(vals) > blm_i and str(vals[blm_i]).strip().lower() == "carrier":
        carrier_ids.add(norm_id(vals[sid_i]))

genes_ws = first_sheet(genes_path)
chip_genes = set()
for row in genes_ws.iter_rows():
    vals = values(row)
    if vals and vals[0] is not None and str(vals[0]).strip():
        chip_genes.add(str(vals[0]).strip().upper())

stage = Counter()
class_counts = Counter()
consequence_counts = Counter()
selected = []
carrier_files = set()
chip_gene_mismatches = 0

for sample_id, path in zip(SAMPLE_IDS, variant_paths):
    if sample_id not in carrier_ids:
        continue
    carrier_files.add(sample_id)
    ws = first_sheet(path)
    headers = values(next(ws.iter_rows(min_row=2, max_row=2)))
    idx = dict((name, index_of(headers, name)) for name in [
        "Chr:Pos", "Zygosity", "Variant Allele Freq", "Gene Names",
        "Sequence Ontology (Combined)", "Classification", "In_CHIP"
    ])
    chip_match_i = idx["In_CHIP"] + 1
    for row in ws.iter_rows(min_row=3):
        vals = values(row)
        stage["carrier_variant_rows"] += 1
        if idx["In_CHIP"] >= len(vals) or not truthy(vals[idx["In_CHIP"]]):
            continue
        stage["chip_annotated"] += 1

        gene_text = vals[idx["Gene Names"]] if idx["Gene Names"] < len(vals) else ""
        match_text = vals[chip_match_i] if chip_match_i < len(vals) else ""
        observed_gene_tokens = set(tokens(gene_text)) | set(tokens(match_text))
        if not (observed_gene_tokens & chip_genes):
            chip_gene_mismatches += 1

        zygosity = str(vals[idx["Zygosity"]] or "").strip()
        if zygosity.lower() == "reference":
            continue
        stage["non_reference"] += 1

        raw_vaf = vals[idx["Variant Allele Freq"]]
        try:
            vaf = float(raw_vaf)
        except (TypeError, ValueError):
            continue
        if not vaf < 0.3:
            continue
        stage["vaf_lt_0_3"] += 1

        consequence = str(vals[idx["Sequence Ontology (Combined)"]] or "").strip()
        low_consequence = consequence.lower()
        excluded_region = ("intron" in low_consequence or
                           "intergenic" in low_consequence or
                           "utr" in low_consequence)
        consequence_counts[(consequence, "excluded" if excluded_region else "retained")] += 1
        if excluded_region:
            continue
        stage["region_retained"] += 1

        classification = str(vals[idx["Classification"]] or "").strip()
        if classification.lower() in ("", "none", "na", "n/a", "-", "not provided"):
            continue
        stage["clinvar_classified"] += 1
        class_counts[classification] += 1

        benign = classification.lower() in ("benign", "likely benign")
        if benign:
            stage["benign_or_likely_benign"] += 1
        selected.append([
            sample_id,
            vals[idx["Chr:Pos"]],
            vals[idx["Gene Names"]],
            consequence,
            zygosity,
            ("%.12g" % vaf),
            classification,
            "1" if benign else "0"
        ])

denominator = stage["clinvar_classified"]
numerator = stage["benign_or_likely_benign"]
if denominator == 0:
    raise RuntimeError("No eligible ClinVar-classified variants; proportion is undefined")
proportion = float(numerator) / float(denominator)

with open("summary.tsv", "w") as out:
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(["metric", "value"])
    w.writerow(["numerator_benign_or_likely_benign", numerator])
    w.writerow(["denominator_all_clinvar_classified", denominator])
    w.writerow(["proportion", repr(proportion)])

with open("selected_variants.tsv", "w") as out:
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(["sample_id", "chr_pos", "gene_names", "sequence_ontology", "zygosity", "vaf", "clinvar_classification", "benign_or_likely_benign"])
    w.writerows(selected)

with open("audit.tsv", "w") as out:
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(["metric", "value"])
    w.writerow(["carrier_status_rows", len(carrier_ids)])
    w.writerow(["carrier_samples_with_variant_files", len(carrier_files)])
    w.writerow(["carrier_samples_missing_variant_file", len(carrier_ids - carrier_files)])
    for key in ["carrier_variant_rows", "chip_annotated", "non_reference", "vaf_lt_0_3", "region_retained", "clinvar_classified", "benign_or_likely_benign"]:
        w.writerow([key, stage[key]])
    w.writerow(["chip_annotation_gene_list_mismatches", chip_gene_mismatches])
    w.writerow(["carrier_sample_ids_with_files", ",".join(sorted(carrier_files))])
    w.writerow(["carrier_sample_ids_missing_files", ",".join(sorted(carrier_ids - carrier_files))])

with open("classifications.tsv", "w") as out:
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(["clinvar_classification", "count"])
    for key in sorted(class_counts):
        w.writerow([key, class_counts[key]])

with open("consequences.tsv", "w") as out:
    w = csv.writer(out, delimiter="\t", lineterminator="\n")
    w.writerow(["sequence_ontology", "region_filter", "count"])
    for key in sorted(consequence_counts):
        w.writerow([key[0], key[1], consequence_counts[key]])

