import csv, io, os, re, sys, tarfile
from collections import Counter
from openpyxl import load_workbook

bundle_path, result_path, audit_path = sys.argv[1:4]
CODING = {
    "synonymous_variant", "missense_variant", "stop_gained", "stop_lost",
    "start_lost", "frameshift_variant", "inframe_insertion", "inframe_deletion",
    "protein_altering_variant", "coding_sequence_variant",
    "incomplete_terminal_codon_variant", "initiator_codon_variant",
    "feature_elongation", "feature_truncation",
}

def values(row):
    return [c.value for c in row]

def norm(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text

def load(tf, member):
    handle = tf.extractfile(member)
    if handle is None:
        raise RuntimeError("Cannot extract " + member.name)
    return load_workbook(io.BytesIO(handle.read()), read_only=False, data_only=True)

def file_sample(name):
    match = re.search(r"CHIP_(.+)\.xlsx$", os.path.basename(name))
    if not match:
        return ""
    sample = match.group(1)
    return sample if sample.startswith("SRR") else sample.split("-", 1)[0]

def so_terms(value):
    if value is None:
        return set()
    return {x.strip() for x in re.split(r"[,;&|]+", str(value)) if x.strip()}

with tarfile.open(bundle_path, "r:*") as tf:
    members = [m for m in tf.getmembers() if m.isfile() and m.name.lower().endswith(".xlsx")]
    cohort_members = [m for m in members if os.path.basename(m.name) == "230215_Trio_Status.xlsx"]
    variant_members = [m for m in members if "CHIP_DP10_GQ20_PASS/" in m.name]
    if (len(members), len(cohort_members), len(variant_members)) != (88, 1, 86):
        raise RuntimeError("Input inventory mismatch")

    wb = load(tf, cohort_members[0])
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows()
    header = values(next(rows))
    sample_index = header.index("Sample ID")
    blm_index = header.index("BLM Mutation Status")
    role_index = header.index("Status")
    cohort = {}
    for row in rows:
        row_values = values(row)
        sample = norm(row_values[sample_index]) if sample_index < len(row_values) else ""
        if sample:
            blm = "" if blm_index >= len(row_values) or row_values[blm_index] is None else str(row_values[blm_index]).strip()
            role = "" if role_index >= len(row_values) or row_values[role_index] is None else str(row_values[role_index]).strip()
            cohort[sample] = (blm, role)

    carriers = {sample for sample, fields in cohort.items() if fields[0].casefold() == "carrier"}
    available = {file_sample(member.name) for member in variant_members}
    evaluable = carriers.intersection(available)
    missing = sorted(carriers - available)
    unmatched = sorted(available - set(cohort))
    if unmatched:
        raise RuntimeError("Variant workbooks absent from cohort table: " + ",".join(unmatched))

    consequence_counts = Counter()
    sample_rows = []
    total_low = total_coding = total_synonymous = 0
    processed = set()
    for member in sorted(variant_members, key=lambda item: item.name):
        sample = file_sample(member.name)
        if sample not in evaluable:
            continue
        processed.add(sample)
        wb = load(tf, member)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows()
        values(next(rows))
        header = values(next(rows))
        try:
            vaf_index = header.index("Variant Allele Freq")
            ontology_index = header.index("Sequence Ontology (Combined)")
        except ValueError as error:
            raise RuntimeError("Required column missing in %s: %s" % (member.name, error))
        sample_low = sample_coding = sample_synonymous = 0
        for row in rows:
            row_values = values(row)
            if max(vaf_index, ontology_index) >= len(row_values):
                continue
            raw_vaf = row_values[vaf_index]
            if raw_vaf is None or str(raw_vaf).strip() == "":
                continue
            try:
                vaf = float(raw_vaf)
            except (TypeError, ValueError):
                continue
            if vaf >= 0.3:
                continue
            sample_low += 1
            total_low += 1
            terms = so_terms(row_values[ontology_index])
            if terms:
                for term in sorted(terms):
                    consequence_counts[term] += 1
            else:
                consequence_counts["<blank>"] += 1
            if terms.intersection(CODING):
                sample_coding += 1
                total_coding += 1
                if "synonymous_variant" in terms:
                    sample_synonymous += 1
                    total_synonymous += 1
        sample_rows.append((sample, cohort[sample][1], cohort[sample][0], sample_low, sample_coding, sample_synonymous))

    if processed != evaluable:
        raise RuntimeError("Not all evaluable carriers were processed")

if total_coding == 0:
    raise RuntimeError("No coding variants with VAF < 0.3")
if total_synonymous > total_coding:
    raise RuntimeError("Numerator exceeds denominator")
fraction = total_synonymous / float(total_coding)

with open(result_path, "w", newline="") as out:
    writer = csv.writer(out, delimiter="\t", lineterminator="\n")
    writer.writerow(["cohort", "vaf_rule", "analysis_unit", "carrier_samples", "coding_variants", "synonymous_variants", "fraction"])
    writer.writerow(["BLM Mutation Status=Carrier with supplied variant workbook", "VAF<0.3", "sample-variant rows", len(processed), total_coding, total_synonymous, format(fraction, ".12g")])

with open(audit_path, "w", newline="") as out:
    writer = csv.writer(out, delimiter="\t", lineterminator="\n")
    writer.writerow(["section", "key", "value1", "value2", "value3", "value4"])
    for key, value in [
        ("xlsx_inputs", len(members)), ("variant_workbooks", len(variant_members)),
        ("cohort_rows", len(cohort)), ("carrier_labels", len(carriers)),
        ("evaluable_carriers", len(evaluable)), ("processed_carrier_samples", len(processed)),
        ("all_low_vaf_rows_in_carriers", total_low), ("coding_low_vaf_rows", total_coding),
        ("synonymous_coding_low_vaf_rows", total_synonymous),
    ]:
        writer.writerow(["scope", key, value, "", "", ""])
    writer.writerow(["scope", "metadata_carriers_without_workbook", ",".join(missing), "", "", ""])
    for sample, role, blm, low, coding, synonymous in sample_rows:
        writer.writerow(["sample", sample, role, blm, low, "%d;%d" % (coding, synonymous)])
    for term, count in sorted(consequence_counts.items()):
        writer.writerow(["consequence", term, count, "coding" if term in CODING else "noncoding", "", ""])
