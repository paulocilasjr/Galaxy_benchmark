import csv
import io
import os
import re
import sys
import tarfile
from collections import Counter
from openpyxl import load_workbook

bundle_path, result_path, audit_path = sys.argv[1:4]
CODING_TERMS = {
    "synonymous_variant", "missense_variant", "stop_gained", "stop_lost",
    "start_lost", "frameshift_variant", "inframe_insertion", "inframe_deletion",
    "protein_altering_variant", "coding_sequence_variant",
    "incomplete_terminal_codon_variant", "initiator_codon_variant",
    "feature_elongation", "feature_truncation",
}

def cell_values(row):
    return [c.value for c in row]

def normalized_sample(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text

def workbook_from_tar(tf, member):
    fh = tf.extractfile(member)
    if fh is None:
        raise RuntimeError("Cannot extract " + member.name)
    return load_workbook(io.BytesIO(fh.read()), read_only=False, data_only=True)

def sample_from_filename(name):
    match = re.search(r"CHIP_(.+)\.xlsx$", os.path.basename(name))
    if not match:
        return ""
    token = match.group(1)
    return token if token.startswith("SRR") else token.split("-", 1)[0]

def so_terms(value):
    if value is None:
        return set()
    return {x.strip() for x in re.split(r"[,;&|]+", str(value)) if x.strip()}

with tarfile.open(bundle_path, "r:*") as tf:
    members = [m for m in tf.getmembers() if m.isfile() and m.name.lower().endswith(".xlsx")]
    cohort_members = [m for m in members if os.path.basename(m.name) == "230215_Trio_Status.xlsx"]
    variant_members = [m for m in members if "CHIP_DP10_GQ20_PASS/" in m.name]
    if len(members) != 88 or len(cohort_members) != 1 or len(variant_members) != 86:
        raise RuntimeError("Input inventory mismatch: total=%d cohort=%d variants=%d" % (len(members), len(cohort_members), len(variant_members)))

    cohort_wb = workbook_from_tar(tf, cohort_members[0])
    cohort_ws = cohort_wb[cohort_wb.sheetnames[0]]
    cohort_rows = cohort_ws.iter_rows()
    cohort_header = cell_values(next(cohort_rows))
    sample_idx = cohort_header.index("Sample ID")
    blm_idx = cohort_header.index("BLM Mutation Status")
    status_idx = cohort_header.index("Status")
    cohort = {}
    for row in cohort_rows:
        vals = cell_values(row)
        sample = normalized_sample(vals[sample_idx]) if sample_idx < len(vals) else ""
        if not sample:
            continue
        blm = "" if blm_idx >= len(vals) or vals[blm_idx] is None else str(vals[blm_idx]).strip()
        role = "" if status_idx >= len(vals) or vals[status_idx] is None else str(vals[status_idx]).strip()
        cohort[sample] = (blm, role)

    carriers = {sample for sample, pair in cohort.items() if pair[0].casefold() == "carrier"}
    consequence_counts = Counter()
    sample_rows = []
    total_low_vaf = 0
    total_coding = 0
    total_synonymous = 0
    processed_carrier_samples = set()

    for member in sorted(variant_members, key=lambda m: m.name):
        sample = sample_from_filename(member.name)
        if sample not in carriers:
            continue
        processed_carrier_samples.add(sample)
        wb = workbook_from_tar(tf, member)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows()
        cell_values(next(rows))
        header = cell_values(next(rows))
        try:
            vaf_idx = header.index("Variant Allele Freq")
            so_idx = header.index("Sequence Ontology (Combined)")
        except ValueError as exc:
            raise RuntimeError("Required column missing in %s: %s" % (member.name, exc))
        sample_low = sample_coding = sample_syn = 0
        for row in rows:
            vals = cell_values(row)
            if max(vaf_idx, so_idx) >= len(vals):
                continue
            raw_vaf = vals[vaf_idx]
            if raw_vaf is None or str(raw_vaf).strip() == "":
                continue
            try:
                vaf = float(raw_vaf)
            except (TypeError, ValueError):
                continue
            if not vaf < 0.3:
                continue
            sample_low += 1
            total_low_vaf += 1
            terms = so_terms(vals[so_idx])
            if not terms:
                consequence_counts["<blank>"] += 1
            else:
                for term in sorted(terms):
                    consequence_counts[term] += 1
            if terms.intersection(CODING_TERMS):
                sample_coding += 1
                total_coding += 1
                if "synonymous_variant" in terms:
                    sample_syn += 1
                    total_synonymous += 1
        sample_rows.append((sample, cohort[sample][1], cohort[sample][0], sample_low, sample_coding, sample_syn))

missing_carrier_files = sorted(carriers - processed_carrier_samples)
unmatched_variant_files = sorted(sample_from_filename(m.name) for m in variant_members if sample_from_filename(m.name) not in cohort)
if missing_carrier_files:
    raise RuntimeError("Carrier samples without variant workbooks: " + ",".join(missing_carrier_files))
if unmatched_variant_files:
    raise RuntimeError("Variant workbooks absent from cohort table: " + ",".join(unmatched_variant_files))
if total_coding == 0:
    raise RuntimeError("No coding variants with VAF < 0.3")
if total_synonymous > total_coding:
    raise RuntimeError("Synonymous numerator exceeds coding denominator")
fraction = total_synonymous / float(total_coding)

with open(result_path, "w", newline="") as out:
    writer = csv.writer(out, delimiter="\t", lineterminator="\n")
    writer.writerow(["cohort", "vaf_rule", "analysis_unit", "carrier_samples", "coding_variants", "synonymous_variants", "fraction"])
    writer.writerow(["BLM Mutation Status=Carrier", "VAF<0.3", "sample-variant rows", len(processed_carrier_samples), total_coding, total_synonymous, format(fraction, ".12g")])

with open(audit_path, "w", newline="") as out:
    writer = csv.writer(out, delimiter="\t", lineterminator="\n")
    writer.writerow(["section", "key", "value1", "value2", "value3", "value4"])
    writer.writerow(["scope", "xlsx_inputs", len(members), "", "", ""])
    writer.writerow(["scope", "variant_workbooks", len(variant_members), "", "", ""])
    writer.writerow(["scope", "cohort_rows", len(cohort), "", "", ""])
    writer.writerow(["scope", "carrier_labels", len(carriers), "", "", ""])
    writer.writerow(["scope", "processed_carrier_samples", len(processed_carrier_samples), "", "", ""])
    writer.writerow(["scope", "all_low_vaf_rows_in_carriers", total_low_vaf, "", "", ""])
    writer.writerow(["scope", "coding_low_vaf_rows", total_coding, "", "", ""])
    writer.writerow(["scope", "synonymous_coding_low_vaf_rows", total_synonymous, "", "", ""])
    for sample, role, blm, low, coding, syn in sample_rows:
        writer.writerow(["sample", sample, role, blm, low, "%d;%d" % (coding, syn)])
    for term, count in sorted(consequence_counts.items()):
        writer.writerow(["consequence", term, count, "coding" if term in CODING_TERMS else "noncoding", "", ""])

