
from __future__ import print_function
import sys, re, os, shutil, tempfile
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

manifest_path = sys.argv[1]
out_path = sys.argv[2]
with open(manifest_path) as handle:
    entries = [line.rstrip('\n').split('\t') for line in handle if line.strip()]
paths = {name: path for name, path in entries}

def workbook_copy_path(path):
    # Galaxy stages datasets with .dat names; openpyxl 2.4 rejects those by extension.
    fd, tmp = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    shutil.copyfile(path, tmp)
    return tmp

def load_xlsx(path):
    tmp = workbook_copy_path(path)
    return load_workbook(tmp, read_only=False, data_only=True)

trio_name = None
chip_name = None
for name in paths:
    if name == '230215_Trio_Status.xlsx':
        trio_name = name
    if name == '230214_Schenz_et_al_2022_CHIP_Genes.xlsx':
        chip_name = name
if trio_name is None:
    raise RuntimeError('Missing trio status workbook')
if chip_name is None:
    raise RuntimeError('Missing CHIP gene workbook')

def norm_sample_id(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if re.match(r'^\d+\.0$', s):
        return s.split('.')[0]
    return s

def rows_from(path):
    wb = load_xlsx(path)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        yield list(row)

trio_rows = list(rows_from(paths[trio_name]))
headers = [str(x).strip() if x is not None else '' for x in trio_rows[0]]
sample_idx = headers.index('Sample ID')
status_idx = headers.index('BLM Mutation Status')
carrier_ids = []
for row in trio_rows[1:]:
    if len(row) <= max(sample_idx, status_idx):
        continue
    if str(row[status_idx]).strip() == 'Carrier':
        sid = norm_sample_id(row[sample_idx])
        if sid:
            carrier_ids.append(sid)
carrier_set = set(carrier_ids)

chip_genes = set()
for row in rows_from(paths[chip_name]):
    if row and row[0] is not None:
        chip_genes.add(str(row[0]).strip())

coding_terms = set([
    'synonymous_variant', 'missense_variant', 'frameshift_variant',
    'inframe_deletion', 'inframe_insertion', 'stop_gained', 'stop_lost',
    'start_lost', 'stop_retained_variant', 'protein_altering_variant',
    'coding_sequence_variant', 'initiator_codon_variant',
    'incomplete_terminal_codon_variant', 'conservative_inframe_deletion',
    'conservative_inframe_insertion', 'disruptive_inframe_deletion',
    'disruptive_inframe_insertion'
])
coding_plus_splice = set(coding_terms)
coding_plus_splice.add('splice_region_variant')
splitter = re.compile(r'[^A-Za-z0-9_]+')

def so_terms(value):
    if value is None:
        return set()
    return set([t for t in splitter.split(str(value).strip()) if t])

def to_decimal(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None

def boolish(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ('true','yes','1'):
        return True
    if s in ('false','no','0'):
        return False
    return None

def sample_from_variant_name(name):
    m = re.search(r'CHIP_(.+)\.xlsx$', name)
    if not m:
        return None
    token = m.group(1)
    if token.startswith('SRR'):
        return token
    return token.split('-')[0]

counts = {
    'all_coding_den': 0, 'all_coding_syn': 0,
    'chip_coding_den': 0, 'chip_coding_syn': 0,
    'chip_coding_splice_den': 0, 'chip_coding_splice_syn': 0,
    'carrier_workbooks': 0, 'variant_rows_seen': 0,
    'vaf_lt_0_3_rows': 0, 'chip_scope_rows': 0,
}
carrier_workbooks = []
missing_required_columns = []
matched_carrier_ids = []

for name, path in paths.items():
    if name in (trio_name, chip_name):
        continue
    sid = sample_from_variant_name(name)
    if sid not in carrier_set:
        continue
    counts['carrier_workbooks'] += 1
    carrier_workbooks.append(name)
    matched_carrier_ids.append(sid)
    wb = load_xlsx(path)
    ws = wb.worksheets[0]
    header_row = None
    header = None
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        labels = [str(x).strip() if x is not None else '' for x in vals]
        if 'Variant Allele Freq' in labels and 'Sequence Ontology (Combined)' in labels:
            header_row = r
            header = labels
            break
    if header_row is None:
        missing_required_columns.append(name)
        continue
    idx = {label: i for i, label in enumerate(header)}
    vaf_i = idx['Variant Allele Freq']
    so_i = idx['Sequence Ontology (Combined)']
    gene_i = idx.get('Gene Names')
    in_chip_i = idx.get('In_CHIP')
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        counts['variant_rows_seen'] += 1
        if len(row) <= max(vaf_i, so_i):
            continue
        vaf = to_decimal(row[vaf_i])
        if vaf is None or not (vaf < Decimal('0.3')):
            continue
        counts['vaf_lt_0_3_rows'] += 1
        terms = so_terms(row[so_i])
        is_syn = 'synonymous_variant' in terms
        is_coding = bool(terms.intersection(coding_terms))
        is_coding_splice = bool(terms.intersection(coding_plus_splice))
        chip_scope = True
        if in_chip_i is not None and len(row) > in_chip_i:
            b = boolish(row[in_chip_i])
            if b is not None:
                chip_scope = b
        elif gene_i is not None and len(row) > gene_i:
            genes = set([g.strip() for g in re.split(r'[,;|]', str(row[gene_i] or '')) if g.strip()])
            chip_scope = bool(genes.intersection(chip_genes))
        if chip_scope:
            counts['chip_scope_rows'] += 1
        if is_coding:
            counts['all_coding_den'] += 1
            if is_syn:
                counts['all_coding_syn'] += 1
        if chip_scope and is_coding:
            counts['chip_coding_den'] += 1
            if is_syn:
                counts['chip_coding_syn'] += 1
        if chip_scope and is_coding_splice:
            counts['chip_coding_splice_den'] += 1
            if is_syn:
                counts['chip_coding_splice_syn'] += 1

if missing_required_columns:
    raise RuntimeError('Missing required variant columns in: ' + ','.join(missing_required_columns))

def frac(num, den):
    if den == 0:
        return 'NA'
    return format(float(Decimal(num) / Decimal(den)), '.12g')

selected_num = counts['chip_coding_syn']
selected_den = counts['chip_coding_den']
selected_fraction = frac(selected_num, selected_den)
with open(out_path, 'w') as out:
    out.write('metric\tvalue\n')
    out.write('selected_fraction_chip_coding_vaf_lt_0_3_synonymous\t%s\n' % selected_fraction)
    out.write('selected_synonymous_count\t%s\n' % selected_num)
    out.write('selected_coding_count\t%s\n' % selected_den)
    for key in sorted(counts):
        out.write('%s\t%s\n' % (key, counts[key]))
    out.write('all_coding_fraction\t%s\n' % frac(counts['all_coding_syn'], counts['all_coding_den']))
    out.write('chip_coding_plus_splice_fraction\t%s\n' % frac(counts['chip_coding_splice_syn'], counts['chip_coding_splice_den']))
    out.write('carrier_ids_from_trio\t%s\n' % ','.join(carrier_ids))
    out.write('matched_carrier_ids\t%s\n' % ','.join(matched_carrier_ids))
    out.write('carrier_workbooks\t%s\n' % ','.join(carrier_workbooks))
