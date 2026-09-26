"""Supplementary Tables, Supplementary Data and the Supplementary Information PDF.

Run from the repository root after build_data.py:  python manuscript_material/scripts/make_supplement.py
Outputs (manuscript_material/supplementary/):
  Supplementary_Tables.xlsx          archive tables renumbered in order of first citation, plus new audit tables
  Supplementary_Table_crosswalk.csv  Supplementary Table number <-> archive ID <-> where the draft cites it
  Supplementary_Data_1..3.xlsx       per-run summaries, failed Galaxy interface calls, parameter substitutions
  Supplementary_Information.pdf      glossary, Supplementary Notes 1-9 and legends for every table and data file
Nothing is regraded, no agent code is executed and no Galaxy server is contacted.
"""
import collections
import csv
import gzip
import json
import os
import re
import statistics as st

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'manuscript_material')
SUP = os.path.join(OUT, 'supplementary')
V2 = os.path.join(ROOT, 'analysis_reports', 'galaxy_improvement_20260924', 'v2_trace_friction')
os.makedirs(SUP, exist_ok=True)

A = json.load(open(os.path.join(ROOT, 'BixBench50_CompBio_analysis', 'analysis.json')))
T = A['tables']
D = json.load(open(os.path.join(OUT, 'source_data', 'figure_data.json')))
LEDGER = json.load(open(os.path.join(V2, 'ledger.json')))
SUMS = [json.loads(l) for l in gzip.open(os.path.join(OUT, 'source_data', 'derived', 'run_summaries.jsonl.gz'), 'rt')]
BENCH = ['BixBench50', 'CompBio', 'IWC']
BLABEL = {'BixBench50': 'BixBench-Verified-50', 'CompBio': 'CompBioBench', 'IWC': 'IWC'}
CFG = {'codex_gpt_5_5': 'GPT-5.5', 'codex_gpt_5_6_sol': 'GPT-5.6 Sol', 'codex_gpt_5_6_luna': 'GPT-5.6 Luna',
       'deepseek_v4_pro_via_codex': 'DeepSeek V4 Pro (Codex)', 'codex_deepseek_v4_pro_0813': 'DeepSeek V4 Pro (Codex)',
       'codex_deepseek_v4_pro': 'DeepSeek V4 Pro (Codex)',
       'deepseek_v4_pro_via_claude_code_superseded': 'DeepSeek V4 Pro (Claude Code, superseded)',
       'codex_gpt_6_astra': 'GPT-6 Astra'}
ENV = {'open_ended_code': 'Open-ended code condition', 'galaxy': 'Galaxy condition'}  # reference condition first everywhere
ENV_ORDER = {'open_ended_code': 0, 'galaxy': 1, 'Open-ended code condition': 0, 'Galaxy condition': 1, 'Open-ended code': 0, 'Galaxy': 1}
CAUSE_NAME = {'SPEC': 'Task under-specified or reference ambiguous', 'RIGOR': 'Statistical or reasoning error',
              'EVALUATOR': 'Evaluator rejected a correct answer', 'KNOWLEDGE': 'Missing domain knowledge',
              'PLATFORM': 'Platform or tool defect', 'HARNESS': 'No answer submitted', 'CONTRACT': 'Output format violated'}
RUN_CFG = {'GPT-5.5': 'GPT-5.5', 'Sol': 'GPT-5.6 Sol', 'Luna': 'GPT-5.6 Luna', 'DS-Codex': 'DeepSeek V4 Pro (Codex)',
           'DS-ClaudeCode': 'DeepSeek V4 Pro (Claude Code, superseded)'}


def load_block(path, start, end):
    """Load constants and functions from an audit script without running its data loaders."""
    src = open(path).read()
    a = src.index(start)
    ns = {'re': re, 'collections': collections, 'json': json}
    exec(src[a:src.index(end, a)], ns)
    return ns


TAXO = load_block(os.path.join(V2, 'taxonomy.py'), 'RULES = [', 'def main')
VAR = load_block(os.path.join(V2, 'variability.py'), 'SH = {', 'def main')
SECRET = re.compile(r'(?i)(api[_-]?key|x-api-key|authorization|password|secret|bearer)(["\'\s:=]+)([^\s"\',}]+)')


def clean(s, limit=2000):
    """Strip Markdown, redact anything shaped like a credential and cap the length for Excel."""
    s = '' if s is None else str(s)
    s = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', lambda m: f'{m.group(1)} ({m.group(2)})' if m.group(2).startswith('http') else m.group(1), s)
    s = s.replace('**', '').replace('`', '')
    s = SECRET.sub(lambda m: m.group(1) + m.group(2) + '[REDACTED]', s)
    s = ILLEGAL_CHARACTERS_RE.sub('', s)
    return s if len(s) <= limit else s[:limit - 1] + '…'


# =====================================================================================================
# New tables produced by this audit (terminology follows scripts/glossary.py)
# =====================================================================================================
def t_ledger():
    rows = []
    for x in sorted(LEDGER, key=lambda x: (x['b'], x['task'], ENV_ORDER[x['cond']], x['run'])):
        _, cfg, rep_ = x['run'].split(' ')
        rows.append({'Benchmark': 'BixBench-Verified-50' if x['b'] == 'BixBench' else 'IWC', 'Task': x['task'],
                     'Execution condition': ENV[x['cond']], 'Model configuration': RUN_CFG[cfg], 'Replicate run': int(rep_[1:]),
                     'Submitted answer or output agreement': clean(x['ans'], 300), 'Decision point that determined the outcome': clean(x['d']),
                     'Primary cause': CAUSE_NAME[x['p']], 'Secondary cause': CAUSE_NAME.get(x['s'], ''), 'Adjudication confidence': x['c']})
    return [('', pd.DataFrame(rows))]


def t_proxy():
    runs = [r for r in A['runs'] if r['benchmark'] == 'CompBio']
    norm = lambda a: (a or '').strip().lower().replace(' ', '')
    tasks = sorted({r['task'] for r in runs})
    modal = {t: collections.Counter(norm(r['answer']) for r in runs if r['task'] == t).most_common(1)[0] for t in tasks}
    per_task = []
    dev = collections.Counter()
    for t in tasks:
        ans, n = modal[t]
        strong = n >= 20
        g = sum(1 for r in runs if r['task'] == t and r['condition'] == 'galaxy' and norm(r['answer']) != ans) if strong else None
        c = sum(1 for r in runs if r['task'] == t and r['condition'] == 'open_ended_code' and r['model'] != 'codex_gpt_6_astra'
                and norm(r['answer']) != ans) if strong else None
        if strong:
            dev['galaxy'] += g
            dev['code'] += c
        per_task.append({'Task': t, 'Runs giving the consensus answer (of 25)': n, 'Strong consensus (20 or more of 25)': 'yes' if strong else 'no',
                         'Open-ended code condition runs deviating (of 12; GPT-6 Astra excluded)': c, 'Galaxy-condition runs deviating (of 12)': g})
    by = collections.defaultdict(dict)
    for r in runs:
        by[(r['model'], r['condition'], r['replicate'])][r['task']] = norm(r['answer'])
    cb = json.load(open(os.path.join(ROOT, 'CompBio', 'compBio_overview_audit.json')))['score_vectors']
    vec = []
    for v in cb:
        if not v.get('answers_compared'):
            continue
        key = (v['model'], v['condition'], int(v['replicate'][1:]))
        vec.append({'Model configuration': CFG.get(v['model'], v['model']), 'Execution condition': ENV[v['condition']], 'Replicate run': int(v['replicate'][1:]),
                    'Reported benchmark score (of 100)': v['score'], 'Archive label': v['score_type'],
                    'Answers matching the consensus answer (of 100)': sum(1 for t, a in by[key].items() if a == modal[t][0])})
    vdf = pd.DataFrame(vec).sort_values('Execution condition', key=lambda c: c.map(ENV_ORDER), kind='stable')
    err = vdf['Answers matching the consensus answer (of 100)'] - vdf['Reported benchmark score (of 100)']
    summary = pd.DataFrame([{'Reported benchmark scores with retained answers': len(vdf), 'Mean absolute error': round(err.abs().mean(), 2),
                             'Mean bias (proxy − reported score)': round(err.mean(), 2),
                             'Strong-consensus tasks': sum(1 for t in tasks if modal[t][1] >= 20),
                             'Open-ended code condition deviations on strong-consensus tasks': dev['code'],
                             'Galaxy-condition deviations on strong-consensus tasks': dev['galaxy']}])
    return [('a | Validation summary', summary), ('b | Proxy against each reported benchmark score', vdf),
            ('c | Consensus support per task', pd.DataFrame(per_task))]


def t_udt():
    s = D['scan']
    us = s['udt_status']
    lab = {'ok': 'Succeeded', 'failed': 'Job failed', 'udt_creation_failed': 'Tool could not be created', 'None': 'No status returned'}
    a = pd.DataFrame([{'Status returned': lab.get(k, k), 'User-defined-tool calls': v, 'Share (%)': round(100 * v / sum(us.values()), 1)}
                      for k, v in sorted(us.items(), key=lambda kv: -kv[1])])
    ph = {'udt': 'User-defined tool', 'tool': 'Installed Galaxy tool', 'wait': 'Job wait call'}
    pn = {'pre_execution_or_command_rendering': 'Before execution or during command rendering', 'command_runtime': 'Command runtime'}
    b = pd.DataFrame([{'Tool type': ph.get(k.split('|')[0], k.split('|')[0]), 'Failure phase': pn.get(k.split('|')[1], k.split('|')[1]),
                       'Error message returned': 'yes' if k.split('|')[2] == 'stderr' else 'no',
                       'Galaxy job errors': v} for k, v in sorted(s['failed_phase'].items(), key=lambda kv: -kv[1])])
    an = s['after_notext']
    c = pd.DataFrame([{'Next tool-run call after a Galaxy job error with no error message': ('Identical request resubmitted' if k.startswith('identical') else 'Request changed'),
                       'Outcome of that call': 'Failed again' if k.endswith('failed') else 'Other outcome', 'Calls': v}
                      for k, v in sorted(an.items())])
    pr = s['probe']
    d = pd.DataFrame([{'Benchmark': BLABEL[bm], 'Galaxy-condition runs with probe tools': pr.get(f'{bm}|runs', 0), 'Probe-tool calls': pr.get(f'{bm}|calls', 0)}
                      for bm in ('BixBench50', 'CompBio')])
    return [('a | Status returned by user-defined-tool calls', a), ('b | Galaxy job errors by tool type, phase and error message', b),
            ('c | What agents did after a Galaxy job error with no error message', c), ('d | Probe tools', d)]


TAX_DEF = {
    'A1': ('Tool description needs an analysis history', 'Tool inspection returned "History unavailable" because no analysis history was specified.'),
    'A2': ('Tool identifier not found', 'A tool identifier that does not resolve on the server (usually guessed or truncated).'),
    'A3': ('Conditional option structure', 'Nested conditional or repeat parameters submitted with a key structure the server rejects.'),
    'A4': ('Parameter value or data type', 'A value, option, format or collection type rejected by tool-state validation.'),
    'A5': ('Dataset or analysis-history identifier', 'Wrong, foreign, truncated or undecodable dataset or analysis-history identifiers, or analysis histories not owned by the user.'),
    'A6': ('User-defined tool definition', 'User-defined tool definition rejected at creation (schema, pattern, discriminator or extra fields).'),
    'A7': ('Upload or file type', 'Unknown file extension or failed upload.'),
    'A8': ('Server, connection or rate limit', 'Web error pages, uncaught server exceptions, closed connections, time-outs and rate limits.'),
    'B1': ('User-defined tool: software missing', 'Package or executable absent from the container chosen for a user-defined tool.'),
    'B2': ('Job failed, no error message', 'A Galaxy job error with no error message returned to the agent.'),
    'B3': ('Job failed with error message', 'A Galaxy job error that returned an error message or traceback.'),
    'B4': ('File format, compression or index', 'Input not parseable in the declared format, compression mismatch or missing index.'),
    'B5': ('Memory or compute limit', 'Memory exhaustion or a killed process.'),
    'Z': ('Unclassified', 'Failure text matched no rule; retained for completeness.'),
}


def t_taxonomy():
    inv = D['inventory']
    rules = {name.split()[0]: rx for name, rx in TAXO['RULES']}
    ex = {}
    for s in SUMS:
        if s['condition'] != 'galaxy':
            continue
        for er in s.get('errors') or []:
            code = TAXO['classify'](er).split()[0]
            if code not in ex and (er.get('err') or '').strip():
                ex[code] = clean(er['err'], 220)
    rows = []
    for code, (name, definition) in TAX_DEF.items():
        cnt = {b: sum(v for k, v in D['ed3b'][b].items() if k.split()[0] == code) for b in BENCH}
        row = {'Class': code, 'Cause': name, 'Stage': 'No job was created' if code.startswith('A') else ('A job was created and failed' if code.startswith('B') else '—'),
               'Definition': definition}
        for b in BENCH:
            row[f'{BLABEL[b]}: failed calls'] = cnt[b]
        row['Total failed calls'] = sum(cnt.values())
        for b in BENCH:
            row[f'{BLABEL[b]}: per 1,000 Galaxy interface calls'] = round(1000 * cnt[b] / inv[b]['mcp_calls'], 2)
        row['Detection rule (regular expression, first match wins)'] = rules.get(code, 'no rule matched')
        row['Example error text'] = ex.get(code, '')
        rows.append(row)
    df = pd.DataFrame(rows)
    tot = {'Class': 'All', 'Cause': 'All failed Galaxy interface calls', 'Stage': '', 'Definition': ''}
    for b in BENCH:
        tot[f'{BLABEL[b]}: failed calls'] = int(df[f'{BLABEL[b]}: failed calls'].sum())
    tot['Total failed calls'] = int(df['Total failed calls'].sum())
    denom = {'Class': 'Denominator', 'Cause': 'All Galaxy interface calls', 'Stage': '', 'Definition': ''}
    for b in BENCH:
        denom[f'{BLABEL[b]}: failed calls'] = inv[b]['mcp_calls']
    denom['Total failed calls'] = sum(inv[b]['mcp_calls'] for b in BENCH)
    df = pd.concat([df, pd.DataFrame([tot, denom])], ignore_index=True)
    sel = lambda p: df[df['Class'].str.fullmatch(p)]
    stage = pd.DataFrame([{'Stage': st_, **{BLABEL[b]: int(sel(p)[f'{BLABEL[b]}: failed calls'].sum()) for b in BENCH}, 'Total': int(sel(p)['Total failed calls'].sum())}
                          for st_, p in [('No job was created (A1–A8)', r'A\d'), ('A job was created and failed (B1–B5)', r'B\d'), ('Unclassified', 'Z')]])
    return [('a | Stage totals', stage), ('b | Causes, definitions, counts and detection rules', df)]


def t_substitution():
    s = D['scan']
    rows = []
    for b in BENCH:
        stt = s['runtool_status'][b]
        tot = sum(stt.values())
        blk, exe = stt.get('validation_parameter_mismatch', 0), stt.get('parameter_mismatch', 0)
        mk = s['mismatch_kind']
        rows.append({'Benchmark': BLABEL[b], 'Tool-run calls (Codex harness)': tot, 'Ran as requested': stt.get('ok', 0),
                     'Parameter substitution blocked before submission': blk, 'Parameter substitution executed': exe,
                     'Any parameter substitution': blk + exe, 'Share of tool-run calls (%)': round(100 * (blk + exe) / tot, 1),
                     'Blocked: value replaced': mk.get(f'{b}|validation_parameter_mismatch|substituted', 0),
                     'Blocked: requested value dropped': mk.get(f'{b}|validation_parameter_mismatch|dropped', 0),
                     'Executed: value replaced': mk.get(f'{b}|parameter_mismatch|substituted', 0),
                     'Executed: requested value dropped': mk.get(f'{b}|parameter_mismatch|dropped', 0)})
    a = pd.DataFrame(rows)
    tools = sorted(s['mismatch_tools'].items(), key=lambda kv: -sum(kv[1].values()))
    b_ = pd.DataFrame([{'Galaxy tool (short identifier)': t, 'Blocked before submission': v.get('validation_parameter_mismatch', 0),
                        'Executed': v.get('parameter_mismatch', 0), 'Total': sum(v.values())} for t, v in tools])
    ex = [('Datamash', 'operations|0|op_name', 'max', 'count'), ('Datamash', 'operations|0|op_column', '6', '1'),
          ('PhyKIT metrics', 'operation|selector', 'treeness', 'total_tree_length'), ('STAR-Fusion', 'genomeDir', 'hg38', 'apiMel4 (honey bee)'),
          ('STAR-Fusion', 'singlePaired|sPaired', 'paired', 'single'), ('Unzip', 'extract_options|target', 'all_regex', 'all'),
          ('DESeq2', 'factorLevel', 'DMSO', 'FactorLevel')]
    c = pd.DataFrame(ex, columns=['Tool', 'Parameter path', 'Value requested by the agent', 'Value Galaxy resolved'])
    return [('a | By benchmark', a), ('b | By tool', b_), ('c | Examples shown in Fig. 3e', c)]


def t_mixed():
    rows = []
    for key, v in D['fig4c_cells'].items():
        b, env, cfg = key.split('|')
        rows.append({'Benchmark': BLABEL[b], 'Execution condition': ENV[env], 'Model configuration': CFG.get(cfg, cfg) if cfg != 'DeepSeek V4 Pro' else 'DeepSeek V4 Pro (Codex)',
                     'Unanimous, all replicate runs succeeded': v.get('all'), 'Split replicate sets': v.get('mixed', 0),
                     'Unanimous, no replicate run succeeded': v.get('none')})
    order = {c: i for i, c in enumerate(['GPT-5.5', 'GPT-5.6 Sol', 'GPT-5.6 Luna', 'DeepSeek V4 Pro (Codex)', 'DeepSeek V4 Pro (Claude Code, superseded)'])}
    bord = {BLABEL[b]: i for i, b in enumerate(BENCH)}
    df = pd.DataFrame(rows).sort_values(['Benchmark', 'Execution condition', 'Model configuration'],
                                        key=lambda c: c.map(order) if c.name == 'Model configuration' else (c.map(ENV_ORDER) if c.name == 'Execution condition' else c.map(bord)))
    tot = df.groupby(['Benchmark', 'Execution condition'], sort=False)['Split replicate sets'].sum().reset_index()
    defs = pd.DataFrame([
        {'Benchmark': 'BixBench-Verified-50', 'Replicate sets per model configuration': 50,
         'Split when': '1 or 2 of 3 replicate runs were scored correct (repeatability category 1–2/3)',
         'Unanimous when': '3/3 or 0/3 replicate runs scored correct'},
        {'Benchmark': 'CompBioBench', 'Replicate sets per model configuration': D['compbio_strong_tasks'],
         'Split when': '1 or 2 of 3 replicate runs matched the consensus answer given by at least 20 of 25 runs (consensus proxy; Supplementary Note 7)',
         'Unanimous when': '3 of 3 or 0 of 3 matched'},
        {'Benchmark': 'IWC', 'Replicate sets per model configuration': '9–10 (tasks scored for that model configuration)',
         'Split when': 'the three output-agreement values ranged by more than 0.05', 'Unanimous when': 'not defined for a continuous endpoint'}])
    return [('a | Definitions', defs), ('b | Totals', tot), ('c | By model configuration', df)]


def t_divergence():
    sh = VAR['SH']
    inv = {v: k for k, v in sh.items()}
    runs = {(s['task'], s['model'], s['condition'], s['replicate']): s for s in SUMS if s['benchmark'] == 'BixBench50'}
    rows = []
    for env, table in (('open_ended_code', VAR['C']), ('galaxy', VAR['G'])):
        for (task, cfg, rep), mech in sorted(table.items()):
            s = runs.get((task, inv[cfg], env, rep), {})
            rows.append({'Execution condition': ENV[env], 'Task': task, 'Model configuration': CFG[inv[cfg]], 'Replicate run': rep,
                         'Divergence mechanism': MECH_NAME[mech], 'Submitted answer': clean(s.get('answer'), 200)})
    df = pd.DataFrame(rows)
    summ = df.groupby(['Divergence mechanism', 'Execution condition']).size().unstack(fill_value=0).reset_index()
    summ = summ[['Divergence mechanism'] + list(ENV.values())]
    for e in ENV.values():
        summ[f'{e} (%)'] = (100 * summ[e] / summ[e].sum()).round(1)
    return [('a | Summary', summ), ('b | Every scored-incorrect replicate run in a split replicate set', df)]


def t_discovery():
    s = D['scan']
    rows = []
    for k, b in enumerate(BENCH):
        sh = D['fig5b'][b]
        n = s['codex_galaxy_runs'][b]
        by = s['bypass'].get(b, {})
        a2 = sum(v for kk, v in D['ed3b'][b].items() if kk.startswith('A2'))
        rows.append({'Benchmark': BLABEL[b], 'Galaxy-condition runs with execution traces (Codex harness)': n,
                     'Median share of Galaxy interface calls spent on tool search and inspection (%)': round(100 * st.median(sh['calls']), 1),
                     'Median share of returned text from tool search and inspection (%)': round(100 * st.median(sh['chars']), 1),
                     'Tool identifiers inspected (summed over runs)': s['inspected'][b],
                     'Inspected but never executed': s['inspected_never_run'][b],
                     'Tool searches per run, median': s['searches'][b]['median'], 'Tool searches per run, maximum': s['searches'][b]['max'],
                     'Unresolvable tool identifiers (cause A2)': a2,
                     'Direct Galaxy API calls among shell commands (%)': round(100 * s['api_shell_core'][b] / s['all_shell'][b], 1),
                     'Runs copying the provided analysis history by direct Galaxy API call': by.get('copy_history', 0) if b == 'BixBench50' else 'not applicable',
                     'Runs downloading outputs by direct Galaxy API call': by.get('download', 0),
                     'Runs reading tool parameter descriptions by direct Galaxy API call': by.get('tool_schema', 0),
                     'Runs checking job status by direct Galaxy API call': by.get('job_polling', 0),
                     'Runs submitting jobs by direct Galaxy API call': by.get('raw_submission', 0)})
    return [('', pd.DataFrame(rows).set_index('Benchmark').T.reset_index().rename(columns={'index': 'Measure'}))]


MECH_NAME = {'V1': 'Galaxy interface trap (silent default, output semantics or job not dispatched)',
             'V3': 'Different software version installed', 'V4': 'Domain convention or definition applied differently',
             'V5': 'Hand-written method instead of the library method (script, or user-defined tool in the Galaxy condition)',
             'V6': 'Error in the final step (sorting, counting, units)', 'V7': 'No answer submitted',
             'V8': 'Answer retrieved from benchmark source files'}

NEW = {
    'NEW-ledger': ('Failure ledger: every scored-incorrect BixBench-Verified-50 run and every IWC run with output agreement below 0.5', t_ledger,
                   'One row per run: 246 scored-incorrect BixBench-Verified-50 runs (135 open-ended code condition, 111 Galaxy condition) and 8 IWC runs '
                   'with output agreement below 0.5. Primary and secondary causes: task under-specified or reference ambiguous (the reference depends '
                   'on a choice the question does not state); statistical or reasoning error (weak evidence relied on, thresholds or invariants ignored, '
                   'wrong denominator); evaluator rejected a correct answer (or the run has a score conflict); missing domain knowledge; platform or '
                   'tool defect (Galaxy, the Galaxy interface or its tools; for open-ended code condition runs, local software); no answer submitted; '
                   'output format violated. Adjudication confidence is defined in Supplementary Note 6.'),
    'NEW-proxy': ('CompBioBench consensus proxy and its validation against reported benchmark scores', t_proxy,
                  'CompBioBench keeps reported benchmark scores but no per-run grades. For each task the most common normalized answer across its 25 '
                  'runs is the consensus answer; tasks where at least 20 of 25 runs agree have a strong consensus. The proxy is validated by counting, '
                  'for each reported benchmark score with retained answers, how many of its 100 answers match the consensus (a). Deviations are '
                  'probable, not adjudicated, failures. Source: CompBio/compBio_overview_audit.json and archived answers.'),
    'NEW-udt': ('Reliability of user-defined tools', t_udt,
                'Galaxy-condition runs with Codex execution traces (1,908 runs), which return a structured status for every tool-run call. Failure '
                'phase and error-message availability come from the failure record of each Galaxy job error. Probe tools are user-defined tools '
                'whose identifier names a probe, preflight, render, smoke, diagnostic or sanity check.'),
    'NEW-taxonomy': ('Causes of every failed Galaxy interface call (A1–A8: no job was created; B1–B5: a job was created and failed)', t_taxonomy,
                     'All Galaxy-condition runs with execution traces, both agent harnesses. Rules are applied in the listed order to the error text '
                     'and status of each failed call; the first match assigns the cause. Rates use all Galaxy interface calls in the benchmark as '
                     'denominator. Example strings are verbatim, truncated and screened for credentials.'),
    'NEW-substitution': ('Parameter substitution in tool-run calls', t_substitution,
                         'The Galaxy interface compares the parameter values an agent requested with the tool state Galaxy resolved from them. A '
                         'parameter substitution is either blocked before submission by the benchmark\'s check (status validation_parameter_mismatch) '
                         'or executed (status parameter_mismatch). "Value replaced" means a requested value was changed; "requested value dropped" '
                         'means a requested parameter is absent from the resolved state. Codex execution traces only. Short tool identifiers are Tool '
                         'Shed repository names.'),
    'NEW-mixed': ('Split replicate sets by benchmark, execution condition and model configuration', t_mixed,
                  'A replicate set is the replicate runs of one task × model configuration × execution condition; it is split when its replicate '
                  'runs disagree in outcome (definitions in a). CompBioBench uses the 82 strong-consensus tasks. The superseded Claude Code model '
                  'configuration was run on BixBench-Verified-50 only.'),
    'NEW-divergence': ('Divergence mechanism of every scored-incorrect replicate run in a split BixBench-Verified-50 replicate set', t_divergence,
                       'Every scored-incorrect replicate run in a split BixBench-Verified-50 replicate set (45 open-ended code condition runs in 33 '
                       'sets; 36 Galaxy-condition runs in 28 sets) was compared call by call with the scored-correct replicate run(s) of the same set. '
                       'The divergence mechanism names what differed between the replicate runs, not why the answer was incorrect (see the primary '
                       'cause in the failure ledger). Definitions in Supplementary Note 8.'),
    'NEW-discovery': ('Tool search and inspection burden, and direct Galaxy API calls', t_discovery,
                      'Galaxy-condition runs with Codex execution traces. Tool search and inspection are search_galaxy_tools and inspect_galaxy_tool '
                      'calls. Returned text is the number of characters returned to the agent, used as a proxy for context consumed because '
                      'input-token usage cannot be attributed to individual calls. Inspected-but-never-executed counts distinct tool identifiers per run '
                      'that were inspected but not submitted, summed over runs. A shell command is a direct Galaxy API call when it uses the BioBlend '
                      'library or calls the Galaxy histories, datasets, jobs or tools API; operations are matched among commands referencing BioBlend '
                      'or any /api/ path. The provided analysis history exists only in BixBench-Verified-50.'),
}

# Order of first citation in the draft Results; NEW-* are tables introduced by this audit.
ORDER = [
    ('X1', 'Design paragraph: benchmarks analysed separately'),
    ('B1', 'Section 1: BixBench accuracy (Fig. 2a)'), ('B10', 'Section 1: task coverage'), ('B11', 'Section 1: tasks scored correct in one condition only'),
    ('C1', 'Section 1: CompBioBench reported benchmark scores (Fig. 2b)'), ('NEW-proxy', 'Section 1: consensus proxy (Extended Data Fig. 1c,d)'),
    ('I1', 'Section 1: IWC mean output agreement (Fig. 2c)'), ('I2', 'Section 1: IWC outcome repeatability'), ('I3', 'Section 1: IWC per task (Fig. 2d)'),
    ('X3', 'Section 1: cross-benchmark tool-set similarity (Extended Data Fig. 2a)'), ('X19', 'Section 1: domain-tool share'),
    ('I4', 'Section 1: IWC condition difference rests on few runs'), ('I6', 'Section 1: IWC score conflicts'),
    ('NEW-ledger', 'Section 1: failure ledger (Fig. 2e)'),
    ('B12', 'Section 1: bix-45-q1 model-configuration dependence (Extended Data Fig. 2b)'), ('B13', 'Section 1: bix-45-q1 version contrast (Extended Data Fig. 2b)'),
    ('B8', 'Section 1: verifier modes, bix-43-q2 (Extended Data Fig. 2c)'), ('I5', 'Section 1: IWC runs below 0.5 (Extended Data Fig. 2d)'),
    ('C4', 'Section 2: biological domains'), ('X14', 'Section 2: workbench operations (Fig. 3a)'),
    ('B9', 'Section 2: most frequent BixBench Galaxy tools (Fig. 3b)'), ('C8', 'Section 2: most frequent CompBioBench Galaxy tools (Fig. 3b)'),
    ('I10', 'Section 2: most frequent IWC Galaxy tools (Fig. 3b)'), ('X12', 'Section 2: user-defined-tool requests and execution (Fig. 3c)'),
    ('X13', 'Section 2: user-defined-tool choice across replicate runs'), ('X15', 'Section 2: why user-defined tools were chosen'),
    ('I11', 'Section 2: IWC user-defined-tool availability'), ('NEW-udt', 'Section 2: user-defined-tool reliability (Extended Data Fig. 3a)'),
    ('B4', 'Section 2: BixBench Galaxy execution evidence'), ('C3', 'Section 2: CompBioBench Galaxy execution evidence'), ('I7', 'Section 2: IWC Galaxy execution evidence'),
    ('X18', 'Section 2: recurring Galaxy job error messages (Fig. 3d)'), ('X5', 'Section 2: error messages retained for Galaxy job errors'),
    ('NEW-taxonomy', 'Section 2: causes of failed Galaxy interface calls (Extended Data Fig. 3b); currently cited as X5'),
    ('X10', 'Section 2: nonzero shell exits'), ('C7', 'Section 2: adjudicated candidate recoveries'),
    ('NEW-substitution', 'Section 2: parameter substitution (Fig. 3e; Extended Data Fig. 3c)'),
    ('X7', 'Section 2: engineering targets (Extended Data Fig. 3d)'),
    ('X9', 'Section 3: software used (Fig. 4a)'), ('I12', 'Section 3: IWC declared routes'),
    ('B5', 'Section 3: BixBench tool-set similarity (Fig. 4b)'), ('C5', 'Section 3: CompBioBench tool-set similarity (Fig. 4b)'),
    ('I8', 'Section 3: IWC tool-set similarity (Fig. 4b)'), ('NEW-mixed', 'Section 3: split replicate sets (Fig. 4c; Extended Data Fig. 4a)'),
    ('X11', 'Section 3: prompt length and workload (Extended Data Fig. 4b)'), ('X16', 'Section 3: different tool-set fingerprints, identical answers'),
    ('B6', 'Section 3: tool-set similarity versus accuracy'), ('NEW-divergence', 'Section 3: divergence mechanisms (Fig. 4d)'),
    ('B7', 'Section 4: BixBench input-token ratio (Fig. 5a)'), ('C6', 'Section 4: CompBioBench input-token ratio'), ('I9', 'Section 4: IWC input-token ratio'),
    ('X2', 'Section 4: input-token ratio across benchmarks (Extended Data Fig. 5a)'),
    ('NEW-discovery', 'Section 4: tool search burden and direct Galaxy API calls (Fig. 5b; Extended Data Fig. 5b)'),
    ('B15', 'Section 4: input-token usage versus accuracy (Fig. 5c)'), ('I13', 'Section 4: IWC model-configuration outcomes'),
]
UNCITED = [('B2', 'Suggest Section 3 (outcome repeatability), beside B5'), ('B3', 'Suggest Section 1 (four near-universal failures; bix-45-q1, bix-30-q3)'),
           ('B14', 'Suggest Section 4 (GPT-5.6 Luna trajectory length)'), ('B16', 'Suggest Section 1 (task coverage sentence)'),
           ('C2', 'Suggest Section 1 or 4 (usage completeness)'), ('C9', 'Suggest Section 3 or 4 (CompBioBench model-configuration outcomes)'),
           ('C10', 'Suggest Methods or limitations (biomedical-knowledge outlier)'), ('X4', 'Suggest Section 4 (run-level burden)'),
           ('X6', 'Suggest Section 2 (where Galaxy job errors concentrate)'), ('X8', 'Suggest Discussion (claim strength)'),
           ('X17', 'Suggest Discussion (six motivating observations)'), ('X20', 'Suggest design paragraph, beside X1')]


def xref(text):
    """Rewrite archive table identifiers inside legends (e.g. 'I3', 'B4/C3/I7') as Supplementary Table numbers."""
    return re.sub(r'\b([BCIX]\d{1,2})\b', lambda m: f'Supplementary Table {TABLE_NO[m.group(1)]}' if m.group(1) in TABLE_NO else m.group(1), text)


# Official terminology (scripts/glossary.py) applied to archive-table titles, headers and legends; cell values stay as archived.
TERM_MAP = [
    (r'Galaxy - open-ended code', 'Condition difference, Galaxy − open-ended code'),
    (r'\bopen-ended-code\b', 'open-ended code'), (r'(?<!open-ended )\bcode runs\b', 'open-ended code runs'),
    (r'\b[Ee]nvironment-wide\b', 'condition-wide'),
    (r'\bEnvironments\b', 'Execution conditions'), (r'\benvironments\b', 'execution conditions'),
    (r'\bEnvironment\b', 'Execution condition'), (r'\benvironment\b', 'execution condition'),
    (r'(?<![Mm]odel )\bConfigurations\b', 'Model configurations'), (r'(?<![Mm]odel )\bconfigurations\b', 'model configurations'),
    (r'(?<![Mm]odel )\bConfiguration\b', 'Model configuration'), (r'(?<![Mm]odel )\bconfiguration\b', 'model configuration'),
    (r'\btriplicate path agreement\b', 'tool-set agreement within replicate sets'), (r'\bmixed cells?\b', 'split replicate sets'),
    (r'\bTriplicates\b', 'Replicate sets'), (r'\btriplicates\b', 'replicate sets'), (r'\bTriplicate\b', 'Replicate set'), (r'\btriplicate\b', 'replicate set'),
    (r'\bWithin-cell\b', 'Within-set'), (r'\bwithin-cell\b', 'within-set'),
    (r'\bCells\b', 'Replicate sets'), (r'\bcells\b', 'replicate sets'), (r'\bcell\b', 'replicate set'),
    (r'\bReplicates\b', 'Replicate runs'), (r'\breplicates\b', 'replicate runs'),
    (r'\bAccepted answers\b', 'Runs scored correct'), (r'\baccepted answers\b', 'runs scored correct'),
    (r'\bAccepted runs\b', 'Runs scored correct'), (r'\baccepted runs\b', 'runs scored correct'),
    (r'\bAccepted\b', 'Scored correct'), (r'\baccepted\b', 'scored correct'),
    (r'\bAcceptance\b', 'Accuracy'), (r'\bacceptance\b', 'accuracy'), (r'\brejected\b', 'scored incorrect'),
    (r'\b([Mm])ean agreement\b', r'\1ean output agreement'), (r'\bAll-task agreement\b', 'All-task output agreement'),
    (r'\bTask-level agreement\b', 'Task-level output agreement'),
    (r'\brecorded paths\b', 'tool-set fingerprints'), (r'\brecorded path\b', 'tool-set fingerprint'), (r'\bpath instrument\b', 'tool-set fingerprint instrument'),
    (r'\bpath agreement\b', 'tool-set agreement'), (r'\bPath agreement\b', 'Tool-set agreement'), (r'\bDifferent recorded paths\b', 'Different tool-set fingerprints'),
    (r'\bError jobs\b', 'Galaxy job errors'), (r'\berror jobs\b', 'Galaxy job errors'), (r'\bGalaxy Galaxy\b', 'Galaxy'),
    (r'\brecovery candidates\b', 'candidate recoveries'), (r'\bCandidate recovery episodes\b', 'Candidate recoveries'),
    (r'\bInput-token burden\b', 'Input-token usage'),
    (r'\banswer-vector scores\b', 'reported benchmark scores'), (r'\baggregate vectors?\b', 'reported benchmark scores'),
    (r'\bscore vectors\b', 'reported benchmark scores'),
    (r'\bMCP\b', 'Galaxy interface'), (r'\bModel dependence\b', 'Model-configuration dependence'), (r'\bModel trade-off\b', 'Model difference'),
    (r'\bnon-fetch Galaxy job errors\b', 'Galaxy job errors'), (r'\bnon-fetch error\b', 'Galaxy job error'),
    (r'\bnon-fetch Galaxy jobs\b', 'Galaxy analysis jobs'), (r'\bnon-fetch jobs\b', 'Galaxy analysis jobs'), (r'\b[Nn]on-fetch\b', 'Galaxy analysis'),
    (r'\bnon-utility Tool Shed jobs\b', 'domain-tool jobs'), (r'\bdomain-analysis\b', 'domain'),
]


OUTPUT_AGREEMENT_TABLES = {'I1', 'I2', 'I3', 'I4', 'I5', 'I6', 'I7', 'I8', 'I9', 'I10', 'I11', 'I12', 'I13', 'X16', 'X20'}


def glossary_terms(text, tid=None):
    for pat, rep in TERM_MAP:
        text = re.sub(pat, rep, text)
    if tid in OUTPUT_AGREEMENT_TABLES:  # IWC agreement is output agreement
        text = re.sub(r'(?<!output )(?<!Output )(?<!tool-set )(?<!Tool-set )\bagreement\b', 'output agreement', text)
        text = re.sub(r'\bAgreement\b(?! with)', 'Output agreement', text)
    elif tid == 'X3':  # path-instrument agreement is tool-set similarity
        text = re.sub(r'\bits agreement\b', 'its tool-set similarity', text).replace('Lower agreement', 'Lower tool-set similarity')
    return text


# =====================================================================================================
# Excel writing
# =====================================================================================================
HEAD = PatternFill('solid', fgColor='E8EEF6')
THIN = Side(style='thin', color='B8B8B8')


def write_sheet(wb, name, number, title, archive_id, blocks, legend):
    ws = wb.create_sheet(name)
    ws['A1'] = f'Supplementary Table {number} | {title}'
    ws['A1'].font = Font(bold=True, size=12)
    ws['A2'] = f'Archive ID: {archive_id}' if not archive_id.startswith('NEW') else 'New table produced by the trace-level audit'
    ws['A2'].font = Font(italic=True, size=9, color='555555')
    r = 4
    widths = collections.defaultdict(int)
    for sub, df in blocks:
        if sub:
            ws.cell(r, 1, sub).font = Font(bold=True, size=10)
            r += 1
        for j, h in enumerate(df.columns, 1):
            c = ws.cell(r, j, clean(h, 300))
            c.font, c.fill, c.border = Font(bold=True, size=9), HEAD, Border(bottom=THIN)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            widths[j] = max(widths[j], min(40, len(str(h)) // 2 + 4))
        r += 1
        for row in df.itertuples(index=False):
            for j, v in enumerate(row, 1):
                v = None if (isinstance(v, float) and pd.isna(v)) else v
                c = ws.cell(r, j, clean(v) if isinstance(v, str) else v)
                c.font = Font(size=9)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                widths[j] = max(widths[j], min(60, len(str(v)) + 2))
            r += 1
        r += 1
    ws.cell(r, 1, 'Notes').font = Font(bold=True, size=9)
    for para in [p for p in clean(legend, 30000).split('\n') if p.strip()]:
        r += 1
        c = ws.cell(r, 1, para)
        c.font = Font(size=9)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(6, min(len(widths), 10)))
        ws.row_dimensions[r].height = max(15, 12 * (len(para) // 140 + 1))
    for j, w in widths.items():
        ws.column_dimensions[get_column_letter(j)].width = max(10, w)
    ws.freeze_panes = 'A5' if not blocks[0][0] else None


def expand_names(x):
    """Spell out benchmark short names used in the archive tables."""
    x = re.sub(r'\bBixBench50\b|\bBix\b', 'BixBench-Verified-50', x)
    return re.sub(r'\bCompBio\b(?!Bench)', 'CompBioBench', x)


def code_first(headers, rows):
    """Put the open-ended code condition before the Galaxy condition wherever an archive table pairs them (reference condition first)."""
    headers, rows = list(headers), [list(r) for r in rows]
    i = 0
    while i < len(headers) - 1:  # adjacent 'Galaxy ...' / 'open-ended code ...' column pairs
        if headers[i].lower().startswith('galaxy') and headers[i + 1].lower().startswith('open-ended code'):
            headers[i], headers[i + 1] = headers[i + 1], headers[i]
            for r in rows:
                r[i], r[i + 1] = r[i + 1], r[i]
            i += 2
        else:
            i += 1
    for j, h in enumerate(headers):  # single columns reported as 'Galaxy / open-ended code'
        if 'Galaxy / open-ended code' in h and all(len(str(r[j]).split(' / ')) == 2 for r in rows):
            headers[j] = h.replace('Galaxy / open-ended code', 'open-ended code / Galaxy')
            for r in rows:
                a, b = str(r[j]).split(' / ')
                r[j] = f'{b} / {a}'
    env = next((j for j, h in enumerate(headers) if h.strip().lower() == 'environment'), None)
    rank = lambda v: 0 if str(v).strip().lower().startswith('open-ended') else (1 if str(v).strip().lower().startswith('galaxy') else None)
    if env is not None:  # within rows sharing the identifying columns before 'Environment', list open-ended code first
        out, i = [], 0
        while i < len(rows):
            j = i
            while j < len(rows) and rows[j][:env] == rows[i][:env]:
                j += 1
            block = rows[i:j]
            if all(rank(r[env]) is not None for r in block):
                block = sorted(block, key=lambda r: rank(r[env]))
            out += block
            i = j
        rows = out
    return headers, rows


def archive_blocks(tid):
    t = T[tid]
    headers, rows = code_first([expand_names(clean(h)) for h in t['headers']], [[expand_names(clean(x)) for x in row] for row in t['rows']])
    headers = [glossary_terms(h[:1].upper() + h[1:], tid) for h in headers]
    return [('', pd.DataFrame(rows, columns=headers))]


def build_tables():
    import glossary
    wb = Workbook()
    idx = wb.active
    idx.title = 'Index'
    entries = [(tid, cite, 'cited') for tid, cite in ORDER] + [(tid, cite, 'not cited in current draft') for tid, cite in UNCITED]
    cross = []
    legends = []
    for n, (tid, cite, status) in enumerate(entries, 1):
        if tid.startswith('NEW'):
            title, fn, legend = NEW[tid]
            blocks = fn()
        else:
            title, legend = glossary_terms(expand_names(clean(T[tid]['title'])), tid), glossary_terms(expand_names(xref(T[tid]['legend'])), tid)
            blocks = archive_blocks(tid)
        write_sheet(wb, f'Supplementary Table {n}', n, title, tid, blocks, legend)
        nrows = sum(len(df) for _, df in blocks)
        cross.append({'Supplementary Table': n, 'Archive ID': tid if not tid.startswith('NEW') else 'new (this audit)',
                      'Internal key': tid, 'Title': title, 'Rows': nrows, 'First cited (draft Results)': cite, 'Status': status})
        legends.append((n, title, legend, tid))
    heads = list(cross[0].keys())
    idx['A1'] = 'Supplementary Tables: index and crosswalk to archive table identifiers'
    idx['A1'].font = Font(bold=True, size=12)
    idx['A2'] = ('Tables are numbered in order of first citation in the draft Results. Tables marked "not cited" should be cited at the '
                 'suggested location or removed before submission; renumber by editing ORDER in scripts/make_supplement.py. Terminology follows '
                 'the Glossary sheet; archived cell values are unchanged.')
    idx['A2'].alignment = Alignment(wrap_text=True)
    idx.merge_cells('A2:G2')
    idx.row_dimensions[2].height = 42
    for j, h in enumerate(heads, 1):
        c = idx.cell(4, j, h)
        c.font, c.fill = Font(bold=True, size=9), HEAD
    for i, row in enumerate(cross, 5):
        for j, h in enumerate(heads, 1):
            c = idx.cell(i, j, row[h])
            c.font = Font(size=9, color='9C3D00' if row['Status'] != 'cited' else '000000')
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for j, w in enumerate([10, 12, 16, 60, 7, 60, 22], 1):
        idx.column_dimensions[get_column_letter(j)].width = w
    idx.freeze_panes = 'A5'
    gl = wb.create_sheet('Glossary', 1)
    gl['A1'] = 'Glossary of official terms used in all figures, tables and supplementary files'
    gl['A1'].font = Font(bold=True, size=12)
    gh = ['Group', 'Official term', 'Status', 'Definition', 'Example in this study', 'Replaces (do not use)']
    for j, h in enumerate(gh, 1):
        c = gl.cell(3, j, h)
        c.font, c.fill = Font(bold=True, size=9), HEAD
    for i, row in enumerate(glossary.rows(), 4):
        for j, v in enumerate(row, 1):
            c = gl.cell(i, j, v)
            c.font = Font(size=9, bold=(j == 2))
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for j, w in enumerate([22, 30, 11, 80, 55, 32], 1):
        gl.column_dimensions[get_column_letter(j)].width = w
    gl.freeze_panes = 'A4'
    wb.save(os.path.join(SUP, 'Supplementary_Tables.xlsx'))
    with open(os.path.join(SUP, 'Supplementary_Table_crosswalk.csv'), 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=heads)
        w.writeheader()
        w.writerows(cross)
    return cross, legends


# =====================================================================================================
# Supplementary Data
# =====================================================================================================
def write_data(fname, title, legend, df):
    wb = Workbook()
    ws = wb.active
    ws.title = 'README'
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=12)
    for i, para in enumerate(legend.split('\n'), 3):
        ws.cell(i, 1, para).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 150
    ws2 = wb.create_sheet('data')
    for j, h in enumerate(df.columns, 1):
        c = ws2.cell(1, j, h)
        c.font, c.fill = Font(bold=True, size=9), HEAD
        ws2.column_dimensions[get_column_letter(j)].width = max(10, min(45, len(h) + 2))
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = None if (isinstance(v, float) and pd.isna(v)) else v
            ws2.cell(i, j, clean(v) if isinstance(v, str) else v)
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = ws2.dimensions
    wb.save(os.path.join(SUP, fname))


def build_data():
    fr = json.load(open(os.path.join(V2, 'per_run_friction.json')))
    rows = []
    for s in sorted(SUMS, key=lambda s: (BENCH.index(s['benchmark']), s['task'], ENV_ORDER[s['condition']], s['model'], s['replicate'])):
        cc = s.get('class_calls') or {}
        f = fr.get(f"{s['benchmark']}|{s['task']}|{s['run_id']}", {}) if s['condition'] == 'galaxy' else {}
        g = s['condition'] == 'galaxy'
        row = {'benchmark': BLABEL[s['benchmark']], 'task': s['task'], 'run_id': s['run_id'], 'model_configuration': CFG.get(s['model'], s['model']),
               'execution_condition': ENV[s['condition']], 'replicate_run': s['replicate'],
               'scored_correct': (s.get('score') == 1) if s['benchmark'] == 'BixBench50' and s.get('score') is not None else None,
               'output_agreement': s.get('score') if s['benchmark'] == 'IWC' else None,
               'submitted_answer': clean(s.get('answer'), 500), 'verifier_mode': s.get('eval_mode') or s.get('verifier'),
               'input_token_usage': s.get('input_tokens'), 'output_tokens': s.get('output_tokens'),
               'execution_trace_present': bool(s.get('trace')), 'shell_commands': s.get('n_shell'), 'nonzero_shell_exits': s.get('n_shell_nonzero'),
               'web_calls': s.get('n_web'), 'galaxy_interface_calls': s.get('n_mcp') if g else None,
               'failed_galaxy_interface_calls': s.get('n_mcp_fail') if g else None,
               'tool_search_and_inspection_calls': cc.get('discovery') if g else None,
               'analysis_history_inspection_calls': cc.get('inspection') if g else None,
               'tool_run_and_job_calls': cc.get('execution') if g else None}
        for code in TAX_DEF:
            row[f'failed_calls_cause_{code}'] = sum(v for k, v in f.items() if k.split()[0] == code) if g else None
        rows.append(row)
    write_data('Supplementary_Data_1_run_summaries.xlsx', 'Supplementary Data 1 | Per-run summary of all 4,240 runs',
               'One row per run (task × model configuration × execution condition × replicate run). scored_correct is the archived BixBench-Verified-50 '
               'evaluator verdict (runs without an answer are scored incorrect); output_agreement is the archived IWC output agreement (0 to 1); '
               'CompBioBench has no per-run grade (its reported benchmark scores are in Supplementary Table 5). No score was regraded.\n'
               'Operational counts come from the execution trace (Codex JSONL or Claude Code stream-JSON). Galaxy interface columns are empty for '
               'open-ended code condition runs. failed_calls_cause_A1..Z use the causes in Supplementary Table '
               f'{TABLE_NO["NEW-taxonomy"]}. Input-token usage includes cached input; output and reasoning tokens are not added.\n'
               'Reference answers are deliberately omitted to limit benchmark contamination; they are available in the source archive evaluator records.',
               pd.DataFrame(rows))
    calls = []
    for s in SUMS:
        if s['condition'] != 'galaxy':
            continue
        for er in s.get('errors') or []:
            calls.append({'benchmark': BLABEL[s['benchmark']], 'task': s['task'], 'run_id': s['run_id'], 'model_configuration': CFG.get(s['model'], s['model']),
                          'trace_line': er.get('line'), 'galaxy_interface_tool': er.get('tool'), 'status': er.get('status'),
                          'cause': TAXO['classify'](er), 'error_text': clean(er.get('err'), 1500), 'arguments': clean(er.get('args'), 1500)})
    old = os.path.join(SUP, 'Supplementary_Data_2_failed_MCP_calls.xlsx')
    if os.path.exists(old):
        os.remove(old)
    write_data('Supplementary_Data_2_failed_Galaxy_interface_calls.xlsx',
               f'Supplementary Data 2 | All {len(calls):,} failed Galaxy interface calls with cause, error text and arguments',
               'One row per failed Galaxy interface call in Galaxy-condition runs with an execution trace (both agent harnesses). trace_line is the '
               'line in the decompressed event log (<benchmark>/analysis/<task>/source_snapshots/huggingface_traces/files/<run_id>/). cause follows '
               f'Supplementary Table {TABLE_NO["NEW-taxonomy"]}. error_text and arguments are verbatim, truncated to 1,500 characters and screened for '
               'credential-like strings.', pd.DataFrame(calls))
    ex = pd.DataFrame(D['scan']['mismatch_examples'], columns=['benchmark', 'task', 'galaxy_tool', 'parameter_path', 'value_requested', 'value_resolved'])
    ex['benchmark'] = ex['benchmark'].map(BLABEL)
    write_data('Supplementary_Data_3_parameter_substitutions.xlsx',
               f'Supplementary Data 3 | {len(ex):,} parameter substitutions with a replaced value (first per tool-run call)',
               'For every tool-run call with a Codex execution trace in which a requested parameter value was replaced (parameter substitution, '
               'blocked or executed), the first replaced parameter with the value requested by the agent and the value Galaxy resolved. Calls in '
               'which a requested value was only dropped are counted in Supplementary Table '
               f'{TABLE_NO["NEW-substitution"]} but have no resolved value to list.', ex)
    return len(calls), len(ex)


# =====================================================================================================
# Supplementary Information PDF
# =====================================================================================================
def build_pdf(legends, n_calls, n_subst):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import glossary

    ss = getSampleStyleSheet()
    body = ParagraphStyle('b', parent=ss['BodyText'], fontName='Helvetica', fontSize=9, leading=12.2, alignment=TA_LEFT, spaceAfter=5)
    h1 = ParagraphStyle('h1', parent=ss['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=17, spaceAfter=8)
    h2 = ParagraphStyle('h2', parent=ss['Heading2'], fontName='Helvetica-Bold', fontSize=11, leading=14, spaceBefore=10, spaceAfter=5)
    small = ParagraphStyle('s', parent=body, fontSize=7.5, leading=9.5, spaceAfter=0)
    smallb = ParagraphStyle('sb', parent=small, fontName='Helvetica-Bold')
    inv = D['inventory']
    fe = D['fig2e']

    def P(t, s=body):
        return Paragraph(t, s)

    def table(data, widths):
        rows = [[P(str(c), smallb if i == 0 else small) for c in r] for i, r in enumerate(data)]
        t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([('LINEABOVE', (0, 0), (-1, 0), 0.6, colors.black), ('LINEBELOW', (0, 0), (-1, 0), 0.4, colors.black),
                               ('LINEBELOW', (0, -1), (-1, -1), 0.6, colors.black), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                               ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
        return t

    S = []
    S += [P('Supplementary Information', h1),
          P('<b>[Manuscript title]</b>'), P('[Authors]'), Spacer(1, 6),
          P('This file contains a glossary of official terms, Supplementary Notes 1–9, legends for Supplementary Tables 1–' + str(len(legends)) +
            ' (provided as Supplementary_Tables.xlsx) and legends for Supplementary Data 1–3 (provided as separate Excel files).'),
          Spacer(1, 6), P('<b>Contents</b>'), P('Glossary of official terms', small)]
    notes = ['Evidence archive and scope', 'Execution-trace extraction, call classes and direct Galaxy API calls', 'Causes of failed Galaxy interface calls',
             'Parameter substitution', 'User-defined tool reliability', 'Failure adjudication protocol', 'CompBioBench consensus proxy',
             'Outcome repeatability: split replicate sets and divergence mechanisms', 'Statistical analysis and limitations']
    for i, n in enumerate(notes, 1):
        S.append(P(f'Supplementary Note {i}. {n}', small))
    S.append(P('Supplementary Table legends; Supplementary Data legends', small))
    S.append(PageBreak())
    # ---- Glossary
    S.append(P('Glossary of official terms', h2))
    S.append(P('These terms are used consistently in every figure, legend, table and data file. Retired wording is listed in the Glossary sheet of '
               'Supplementary_Tables.xlsx.', body))
    for g, rr in glossary.GROUPS:
        S.append(P(f'<b>{g}</b>', body))
        S.append(table([['Official term', 'Definition', 'Example in this study']] + [[t, d, e] for t, s_, d, e, r in rr], [34, 86, 50]))
        S.append(Spacer(1, 6))
    S.append(PageBreak())
    # ---- Note 1
    S += [P('Supplementary Note 1. Evidence archive and scope', h2),
          P(f'The analysis is retrospective and read-only. It uses 4,240 runs: BixBench-Verified-50 (50 tasks; {inv["BixBench50"]["runs"]:,} runs: five model '
            'configurations × two execution conditions × three replicate runs), CompBioBench (100 tasks; 2,400 paired runs from four model configurations, '
            'plus 100 unpaired GPT-6 Astra runs in the open-ended code condition) and IWC (10 tasks; 240 runs, four model configurations). Of these, '
            f'{sum(inv[b]["traces"] for b in BENCH):,} retain an execution trace (Codex JSONL, or Claude Code stream-JSON for the superseded DeepSeek V4 Pro '
            'model configuration); the 12 missing traces are CompBioBench runs. Galaxy job records comprise '
            f'{sum(inv[b]["nonfetch_jobs"] for b in BENCH):,} Galaxy analysis jobs ({inv["BixBench50"]["nonfetch_jobs"]:,} BixBench, '
            f'{inv["CompBio"]["nonfetch_jobs"]:,} CompBioBench, {inv["IWC"]["nonfetch_jobs"]:,} IWC), counted once per server and job identifier; '
            'data-fetch jobs are excluded. Detailed analysis-history snapshots exist for '
            f'{inv["BixBench50"]["detailed_histories"][0]} of {inv["BixBench50"]["detailed_histories"][1]}, '
            f'{inv["CompBio"]["detailed_histories"][0]:,} of {inv["CompBio"]["detailed_histories"][1]:,} and '
            f'{inv["IWC"]["detailed_histories"][0]} of {inv["IWC"]["detailed_histories"][1]} Galaxy-condition runs; they were taken after the runs completed.'),
          P('Archive integrity was checked before analysis: 160 source-evidence hashes were verified, all 4,240 run identifiers are unique, and the '
            'archived accuracy, Galaxy job states and input-token medians were reproduced from the raw records. No score was regraded; the original '
            'evaluator verdict of every run is retained, and runs without an answer count as scored incorrect. No agent code was executed and no Galaxy '
            'server was contacted. BixBench reference values were read only from the evaluator records of completed runs.'),
          P('The three benchmarks keep different endpoints: accuracy (BixBench-Verified-50, binary evaluator verdicts), reported benchmark scores without '
            'per-run grades (CompBioBench) and continuous output agreement with workflow reference outputs (IWC). They are analysed separately and never '
            'pooled. Extended Data Fig. 1a summarizes the archive and Extended Data Fig. 1b the audit pipeline.')]
    # ---- Note 2
    s = D['scan']
    S += [P('Supplementary Note 2. Execution-trace extraction, call classes and direct Galaxy API calls', h2),
          P('Every execution trace was parsed. For each Galaxy interface call we recovered the tool name, arguments, structured status, error text, '
            'failure summary, validation errors and parameter provenance. Calls were grouped as <i>tool search and inspection</i> (search_galaxy_tools, '
            'inspect_galaxy_tool), <i>analysis-history inspection</i> (history and archive inspection), <i>tool runs and jobs</i> (run_galaxy_tool_and_wait, '
            'run_galaxy_udt_and_wait, job waits) and shell commands. Input-token usage is recorded per run and cannot be attributed to individual calls, '
            'so the number of characters each call returned to the agent is used as a proxy for context consumed (Fig. 5b).'),
          P('A shell command in a Galaxy-condition run is a direct Galaxy API call when it uses the BioBlend library or calls the Galaxy histories, '
            'datasets, jobs or tools API. Commands referencing BioBlend or any /api/ path were matched against five operations: copying the provided '
            'analysis history (copy_history, /copy), downloading datasets (download_dataset, /display, /download), reading tool parameter descriptions '
            '(show_tool, /build), checking job status (show_job, /jobs/{id}, wait_for_job) and submitting jobs directly (run_tool, or a POST to '
            '/api/tools through requests, curl or urllib). Direct Galaxy API calls, parameter substitution and user-defined-tool statistics use Codex '
            f'execution traces ({sum(s["codex_galaxy_runs"].values()):,} Galaxy-condition runs: {s["codex_galaxy_runs"]["BixBench50"]} BixBench, '
            f'{s["codex_galaxy_runs"]["CompBio"]:,} CompBioBench, {s["codex_galaxy_runs"]["IWC"]} IWC), because the Claude Code agent harness records tool '
            'results in a different structure. A provided analysis history exists, and copying it is required by the harness policy, only in BixBench.')]
    # ---- Note 3
    tax_rows = [['Class', 'Cause', 'Definition']] + [[k, v[0], v[1]] for k, v in TAX_DEF.items()]
    tot = {b: sum(D['ed3b'][b].values()) for b in BENCH}
    pre = sum(v for b in BENCH for k, v in D['ed3b'][b].items() if k.startswith('A'))
    post = sum(v for b in BENCH for k, v in D['ed3b'][b].items() if k.startswith('B'))
    unc = sum(v for b in BENCH for k, v in D['ed3b'][b].items() if k.startswith('Z'))
    S += [P('Supplementary Note 3. Causes of failed Galaxy interface calls', h2),
          P(f'Of {sum(inv[b]["mcp_calls"] for b in BENCH):,} Galaxy interface calls, {sum(tot.values()):,} returned a failure status. Each failed call was '
            'assigned one cause by applying ordered regular-expression rules to its error text and status (first match wins; rules in Supplementary '
            f'Table {TABLE_NO["NEW-taxonomy"]}). Causes A1–A8 occurred before any job was created ({pre:,} calls); B1–B5 are failures of a created job, '
            f'that is Galaxy job errors ({post:,}); {unc} calls matched no rule and are reported as unclassified. Rule order resolves overlaps; for '
            'example, a missing analysis history is assigned A1 even when the message also contains a validation keyword. Each cause maps onto an '
            'interface change (Extended Data Fig. 3d).'),
          table(tax_rows, [12, 50, 108])]
    # ---- Note 4
    S += [P('Supplementary Note 4. Parameter substitution', h2),
          P('The Galaxy interface compares the parameter values an agent requested with the tool state Galaxy resolved from them, and returns the '
            'difference as parameter provenance. A parameter substitution is reported with one of two statuses: <i>validation_parameter_mismatch</i>, '
            'when the benchmark\'s check blocked submission, and <i>parameter_mismatch</i>, when the job executed. A value is <i>replaced</i> when a '
            'requested value was changed (for example max → count) and <i>dropped</i> when a requested parameter is absent from the resolved state, '
            'typically because it was placed under a conditional option or repeat that Galaxy did not select. Parameter substitution occurred in 4,352 '
            f'tool-run calls: 3,436 blocked and 916 executed (Fig. 3e; Supplementary Table {TABLE_NO["NEW-substitution"]}). Supplementary Data 3 lists '
            f'the first replaced value of each call with a replaced value ({n_subst:,} calls). Whether a substitution changed a scientific result was '
            'established only for the adjudicated runs in the failure ledger.')]
    # ---- Note 5
    S += [P('Supplementary Note 5. User-defined tool reliability', h2),
          P('Every run_galaxy_udt_and_wait call returns a structured status (ok, failed, udt_creation_failed or none). For each Galaxy job error, the '
            'failure record gives the phase (before execution or during command rendering, versus command runtime) and whether an error message was '
            'returned. A Galaxy job error with no error message was followed to the next tool-run call of the same run, which was classed as an '
            'identical request when its tool inputs or user-defined-tool definition were byte-identical after key sorting, and otherwise as changed. '
            'Probe tools were identified by identifiers containing probe, preflight, render, smoke, diagnos(tic) or sanity. Of the calls that followed a '
            'Galaxy job error with no error message, 784 resubmitted an identical request and 615 of those failed again (Extended Data Fig. 3a; '
            f'Supplementary Table {TABLE_NO["NEW-udt"]}).')]
    # ---- Note 6
    cat_rows = [['Cause', 'Definition', 'Open-ended code condition, primary / secondary (n = 135)', 'Galaxy condition, primary / secondary (n = 111)']]
    cdef = {'SPEC': 'The question is under-specified, or the reference depends on a choice the question does not state',
            'RIGOR': 'Weak evidence relied on, thresholds or invariants ignored, or a wrong denominator',
            'EVALUATOR': 'The evaluator rejected an equivalent answer, or the run has a score conflict',
            'KNOWLEDGE': 'A domain-knowledge gap', 'PLATFORM': 'A defect of Galaxy, the Galaxy interface or its tools (for open-ended code condition runs: local software)',
            'HARNESS': 'No answer written (turn ended or budget exhausted)', 'CONTRACT': 'Output-format or identifier violation'}
    for k, v in cdef.items():
        cat_rows.append([CAUSE_NAME[k], v, f'{fe["open_ended_code"]["primary"].get(k, 0)} / {fe["open_ended_code"]["secondary"].get(k, 0)}',
                         f'{fe["galaxy"]["primary"].get(k, 0)} / {fe["galaxy"]["secondary"].get(k, 0)}'])
    S += [P('Supplementary Note 6. Failure adjudication protocol', h2),
          P('Every scored-incorrect BixBench-Verified-50 run (246: 135 open-ended code condition, 111 Galaxy condition) and every IWC run with output '
            'agreement below 0.5 (8) was read call by call. For each, we identified the decision point: the step at which the evidence the agent acted on '
            'diverged from what the reference required, and whether that evidence was sufficient. Scored-incorrect runs were contrasted with '
            'scored-correct runs of the same task, using the same model configuration where one succeeded. Final answers were checked against independent '
            'evidence where possible (for example, reconstructing a statistic from the retained inputs, recomputing Benjamini–Hochberg adjustments, or '
            'comparing assembled contigs with the public <i>Agrius convolvuli</i> mitochondrial genome OZ203683.1 by canonical 31-base sequences) rather '
            'than against the agent\'s own conclusion.'),
          P('Each run received one primary cause and at most one secondary cause. Galaxy is counted as <i>implicated</i> when a platform or tool defect is '
            f'the primary or secondary cause: {fe["open_ended_code"]["platform_any"]} of 135 scored-incorrect open-ended code condition runs and '
            f'{fe["galaxy"]["platform_any"]} of 111 scored-incorrect Galaxy-condition runs.'),
          table(cat_rows, [40, 70, 30, 30]), Spacer(1, 5),
          P('Adjudication confidence was assigned per run: <b>high</b> (159), the decisive step is visible in the execution trace and confirmed by '
            'reproduction, independent reference comparison or contrast with a scored-correct sibling; <b>moderate</b> (55), the decisive step is '
            'identified but only partly verified; <b>mixed</b> (20), several contributing causes are present and the primary assignment is a judgement; '
            '<b>unresolved</b> (12), no decisive step could be isolated and the cause is the best-supported explanation. Mixed and unresolved assignments '
            '(32 of 246) are not proofs. The failure ledger (Supplementary Table ' + str(TABLE_NO['NEW-ledger']) + ') records the decision point for every '
            'run. For IWC, the two Galaxy-condition host-removal runs with a score conflict (evaluator 0.273, run record 0.9999 and 1.0) were '
            're-examined: the submitted BWA-MEM output kept 20,899 read pairs, matching the BWA-route reference (20,896) rather than the Bowtie2-route '
            'reference (72,867) against which it was scored; the IWC archive itself does not adjudicate these score conflicts (Extended Data Fig. 2d).')]
    # ---- Note 7
    S += [P('Supplementary Note 7. CompBioBench consensus proxy', h2),
          P('CompBioBench archives 22 reported benchmark scores but no per-run grades or references, so per-task correctness cannot be observed. '
            'Answers were normalized (lower case, whitespace removed) and the most common answer across the 25 runs of each task was taken as the '
            'consensus answer. For the 22 reported benchmark scores with retained answers, the number of answers matching the consensus reproduces the '
            'reported score with a mean absolute error of 2.6 of 100 and a mean bias of +2.0 (Extended Data Fig. 1c). On the 82 tasks where at least 20 '
            'of 25 runs agree, a deviating answer is treated as a <i>probable</i> failure; this yields 40 open-ended code condition and 33 Galaxy-condition '
            'deviations (GPT-6 Astra excluded). The 18 contested tasks are not used. Re-fitting individual answers to the reported scores was rejected '
            'because it over-fits (22 constraints, 100 unknowns). The consensus proxy is labelled as such wherever it appears and supports no per-task '
            'accuracy claim.')]
    # ---- Note 8
    S += [P('Supplementary Note 8. Outcome repeatability: split replicate sets and divergence mechanisms', h2),
          P('A replicate set is the replicate runs of one task × model configuration × execution condition (the "cell" of the data table). A replicate set '
            'is <i>split</i> when its replicate runs disagree in outcome: 1 or 2 of 3 BixBench-Verified-50 replicate runs scored correct (repeatability '
            'category 1–2/3); 1 or 2 of 3 CompBioBench replicate runs match the consensus answer (82 strong-consensus tasks); or the within-set range of '
            'IWC output agreement exceeds 0.05. For every scored-incorrect replicate run in a split BixBench-Verified-50 replicate set, the execution trace '
            'was compared with those of the scored-correct sibling(s) to identify the divergence mechanism: <b>Galaxy interface trap</b>, an optional '
            'interface path taken by this replicate run only produced a silent default, a different output semantics or an undispatched job; '
            '<b>different software version installed</b>; <b>domain convention or definition applied differently</b>; <b>hand-written method instead of '
            'the library method</b> (a script, or a user-defined tool in the Galaxy condition); <b>error in the final step</b> (sorting, counting, units) '
            'after an otherwise correct analysis; <b>no answer submitted</b>; and <b>answer retrieved from benchmark source files</b>, when the only '
            'scored-correct replicate run retrieved benchmark source data. A divergence mechanism describes what differed between replicate runs and is '
            'distinct from the primary cause (Note 6). Replicate labels are not matched random seeds, so outcome repeatability means run-to-run agreement '
            'under the same prompt and model configuration.')]
    # ---- Note 9
    S += [P('Supplementary Note 9. Statistical analysis and limitations', h2),
          P('Unless stated otherwise, intervals are exploratory 95% percentile cluster-bootstrap intervals with 20,000 resamples (seed 20260922; '
            'NumPy linear quantiles; statistic-specific SHA-256 streams). Resampling units are source capsules for BixBench (up to 33), tasks for '
            'CompBioBench (up to 100) and tasks for IWC; model configurations and replicate runs are retained within each cluster, and cross-benchmark '
            'draws are independent and stratified by benchmark. Intervals assume independent clusters, which is not established for shared biological '
            'inputs, and are pointwise rather than multiplicity-adjusted. No confirmatory significance, equivalence, non-inferiority or causal claim is '
            'made; an interval spanning zero (differences) or one (ratios) is inconclusive. The input-token ratio divides the median input-token usage of '
            'the three Galaxy-condition runs by the median of the three open-ended code condition runs for the same task and model configuration (all six '
            'values required), and is summarized by its median. Spearman correlations between prompt length or recorded workload and operational errors '
            'are across tasks, use the three shared GPT model configurations (nine observed runs required per task and execution condition) and 5,000 '
            'bootstrap resamples with ties re-ranked within each resample.'),
          P('<b>Limitations.</b> Benchmarks differ in task selection, prompts, execution budgets, exposed interfaces and endpoints, so cross-benchmark '
            'contrasts are descriptive. The IWC estimates rest on nine or ten task clusters. The superseded Claude Code model configuration ran '
            'BixBench-Verified-50 only. Analysis histories were snapshotted after the runs. Adjudication was performed by execution-trace review, and 32 '
            'of 246 assignments are mixed or unresolved. Candidate recoveries were not adjudicated, and the open-ended code condition has no equivalent '
            'job-level record. Nonzero shell exits are an operational marker, not a scientific failure rate. Human review time, reconstruction accuracy '
            'and monetary cost were not measured.')]
    S.append(PageBreak())
    # ---- Table legends
    S.append(P('Supplementary Table legends', h2))
    S.append(P('All tables are provided in Supplementary_Tables.xlsx, one sheet per table, with an index sheet mapping each table to its archive '
               'identifier and first citation, and a glossary sheet.', body))
    for n, title, legend, tid in legends:
        leg = clean(legend, 4000).replace('\n\n', ' ').replace('\n', ' ')
        leg = re.sub(r'Unless a table specifies otherwise, new intervals are exploratory.*', 'Intervals as in Supplementary Note 9.', leg)
        src = f' Archive table {tid}.' if not tid.startswith('NEW') else ''
        ttl = clean(title) if clean(title).endswith(('?', '.')) else clean(title) + '.'
        S.append(P(f'<b>Supplementary Table {n} | {ttl}</b> {leg}{src}', small))
        S.append(Spacer(1, 4))
    S.append(P('Supplementary Data legends', h2))
    S += [P('<b>Supplementary Data 1 | Per-run summary of all 4,240 runs.</b> One row per run with benchmark, task, model configuration, execution '
            'condition, replicate run, evaluator verdict or output agreement, submitted answer, verifier mode, input-token usage, shell and Galaxy '
            'interface call counts and failed-call counts by cause. Reference answers are omitted to limit benchmark contamination.', small), Spacer(1, 4),
          P(f'<b>Supplementary Data 2 | All {n_calls:,} failed Galaxy interface calls.</b> One row per failed call with run, trace line, interface tool, '
            'status, cause, verbatim error text and arguments (truncated to 1,500 characters; screened for credentials).', small), Spacer(1, 4),
          P(f'<b>Supplementary Data 3 | {n_subst:,} parameter substitutions with a replaced value.</b> For each tool-run call with a Codex execution trace '
            'in which a requested value was replaced, the first replaced parameter with the value requested and the value Galaxy resolved.', small)]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 7.5)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f'Supplementary Information — page {doc.page}')
        canvas.restoreState()

    doc = SimpleDocTemplate(os.path.join(SUP, 'Supplementary_Information.pdf'), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm, title='Supplementary Information', author='[Authors]')
    doc.build(S, onFirstPage=footer, onLaterPages=footer)


TABLE_NO = {tid: n for n, (tid, _) in enumerate(ORDER + UNCITED, 1)}

if __name__ == '__main__':
    cross, legends = build_tables()
    n_calls, n_subst = build_data()
    build_pdf(legends, n_calls, n_subst)
    print(f'{len(cross)} Supplementary Tables ({sum(1 for c in cross if c["Status"] == "cited")} cited); '
          f'Supplementary Data: {n_calls} failed calls, {n_subst} substitutions')
