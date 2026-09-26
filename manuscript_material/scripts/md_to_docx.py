"""Convert the Markdown legends and statements to Word files for pasting into the manuscript.

Run from the repository root:  python manuscript_material/scripts/md_to_docx.py
Handles headings (#, ##), paragraphs, '- ' bullets, **bold** and *italic*; that is all these files use.
"""
import os
import re

from docx import Document
from docx.shared import Pt

OUT = os.path.join(os.path.dirname(__file__), '..')
FILES = ['legends/Figure_legends.md', 'statements/Data_availability.md', 'statements/Code_availability.md',
         'statements/Reporting_summary_notes.md']


def add_runs(par, text):
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not tok:
            continue
        if tok.startswith('**'):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith('*'):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok.replace('`', ''))


def convert(src):
    """Legends become one paragraph per figure (Nature style: title, then run-in panels); other files keep their paragraphs."""
    legends = 'legend' in src.lower()
    doc = Document()
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Arial', Pt(10)
    para = []

    def flush():
        if para:
            add_runs(doc.add_paragraph(), ' '.join(para))
            para.clear()

    for line in open(src, encoding='utf-8'):
        line = line.rstrip('\n')
        if line.startswith('#'):
            flush()
            level = len(line) - len(line.lstrip('#'))
            doc.add_heading(line.lstrip('# ').strip(), level=min(level, 3))
        elif legends:
            if re.match(r'^\*\*(Extended Data )?Fig\. \d', line):
                flush()
            if line.strip():
                para.append(line[2:] if line.startswith('- ') else line)
        elif line.startswith('- '):
            flush()
            add_runs(doc.add_paragraph(style='List Bullet'), line[2:])
        elif not line.strip():
            flush()
        else:
            para.append(line)
    flush()
    dst = os.path.splitext(src)[0] + '.docx'
    doc.save(dst)
    return dst


if __name__ == '__main__':
    for f in FILES:
        p = os.path.join(OUT, f)
        if os.path.exists(p):
            print('wrote', os.path.relpath(convert(p), OUT))
