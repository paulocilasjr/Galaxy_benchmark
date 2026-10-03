"""Check the integrated draft and preserve a local source/output hash manifest.

This validates archive accounting and artifact contents. It does not establish
scientific correctness, native application compatibility or submission readiness.
"""
from pathlib import Path
import hashlib
import gzip
import importlib.metadata
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
import pandas as pd
import openpyxl
from pypdf import PdfReader

PAPER=Path(__file__).resolve().parents[1]
NARR=PAPER.parent
ROOT=NARR.parent

def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''):h.update(chunk)
    return h.hexdigest()

def main():
    errors=[]
    def check(condition,message):
        if not condition:errors.append(message)
    md=(PAPER/'manuscript.md').read_text()
    numbers=json.loads((PAPER/'numbers.json').read_text())
    prov=json.loads((PAPER/'numbers_provenance.json').read_text())
    refs=json.loads((NARR/'references.json').read_text())
    for key in re.findall(r'\{\{([\w.-]+)\}\}',md):check(key in numbers and key in prov,'Missing numeric value/source: '+key)
    for group in re.findall(r'\[@([^\]]+)\]',md):
        for key in group.split(';'):check(key.strip().lstrip('@') in refs,'Unknown reference: '+key)
    results=md.split('## Results\n',1)[1].split('## Discussion',1)[0]
    check(len(re.findall(r'^### [1-4]\. ',results,re.M))==4,'Four requested Results sections are required')
    check('results pending' in md.lower() and 'RESULTS PENDING' in (PAPER/'TOKEN_REDUCTION_PLACEHOLDER.md').read_text().upper(),'Prospective status is absent')
    author=json.loads((NARR/'author_metadata.json').read_text())
    pending=[k.strip() for k in re.findall(r'\[Authors:\s*([^\]]+)\]',md) if not author['fields'].get(k.strip())]
    runs=pd.read_csv(PAPER/'analysis/accuracy_primary_runs.csv')
    sets=pd.read_csv(PAPER/'analysis/accuracy_replicate_sets.csv')
    check(len(runs)==3816 and len(sets)==1272,'Primary run/set accounting changed')
    check(runs.groupby(['benchmark','task','cfg','env']).size().eq(3).all(),'A primary replicate set is not complete')
    check(not runs.duplicated(['benchmark','task','cfg','env','replicate']).any(),'Duplicate primary run keys')
    check(sets.n.eq(3).all(),'Replicate sets must contain three scored attempts')
    iwc=sets[sets.benchmark.eq('IWC')]
    check(np.array_equal(iwc.discordant.to_numpy(),iwc.within_set_range.gt(.05).astype(int).to_numpy()),'IWC discordance differs from declared >0.05 rule')
    cases=pd.read_csv(PAPER/'analysis/rigor_task_cases.csv')
    check(len(cases)==93,'Audited case accounting changed')
    binary=sets[sets.benchmark.isin(['BixBench50','CompBio'])]
    failing={({'BixBench50':'BixBench-Verified-50','CompBio':'CompBioBench'}[b],t) for (b,t),e in binary.groupby(['benchmark','task']).n_errors.max().items() if e>0}
    audited=set(zip(cases.benchmark,cases.task))
    check(failing<=audited,'A task with a rejected primary run is missing from the audit census')
    check(int(cases.primary_failing_task.sum())==len(failing)==73,'Census of failing binary-benchmark tasks changed')
    totals=pd.read_csv(PAPER/'analysis/token_arm_total_ratios.csv')
    check(np.allclose(totals.ratio_of_totals,totals.galaxy_total_tokens/totals.code_total_tokens),'Ratio of totals does not equal total Galaxy / total code tokens')
    fc=pd.read_csv(PAPER/'analysis/token_failure_classes.csv')
    for population in ['four_primary_configurations','all_archived_configurations']:
        ops=pd.read_csv(PAPER/'analysis/token_operation_volume.csv')
        ops=ops[(ops.population==population)&(ops.benchmark=='All benchmarks')]
        check(fc[fc.population==population].calls.sum()==ops.failed_calls.sum(),'Failure classes do not reconcile with the failure union: '+population)
    for term in ['TOKEN_REDUCTION_PLACEHOLDER','manuscript_narrative/','parent package','Contributor discussion','.csv','.json']:
        check(term not in md,'Manuscript contains an internal file or package reference: '+term)
    op=pd.read_csv(PAPER/'analysis/token_operation_volume.csv')
    p=op[(op.population=='four_primary_configurations')&(op.benchmark=='All benchmarks')]
    check(p.calls.sum()==66316,'Primary interface call count changed')
    allp=op[(op.population=='all_archived_configurations')&(op.benchmark=='All benchmarks')]
    check(allp.calls.sum()==69505 and allp.failed_calls.sum()==7987,'All-archive interface union changed')
    check(p[p.operation=='run_galaxy_udt_and_wait'].calls.sum()==4960,'Primary UDT call count changed')
    cells=pd.read_csv(PAPER/'analysis/token_cell_eligibility_and_ratios.csv')
    valid=cells[cells.eligible]
    check(np.allclose(valid.ratio,valid.galaxy/valid.open_ended_code),'Token ratio numerator/denominator changed')
    check((valid.open_ended_code>0).all(),'Eligible token denominator is not positive')
    pdfs=sorted((PAPER/'figures').glob('Fig*.pdf'))
    check([p.stem for p in pdfs]==[f'Fig{i}' for i in range(1,7)],'Six main PDFs required')
    for p in pdfs:
        r=PdfReader(p);check(len(r.pages)==1,'Figure is not one page: '+p.name)
        check(float(r.pages[0].mediabox.width)*25.4/72<=180.01,'Figure exceeds180mm: '+p.name)
    payload=json.loads((PAPER/'source_data/figure_source_data.json').read_text())
    docs={(s,c) for s,c,d in payload['dictionary'] if d and 'Value shown in the figure' not in d}
    wb=openpyxl.load_workbook(PAPER/'Source_Data.xlsx',read_only=True,data_only=True)
    check(set(wb.sheetnames)=={'README','Column_dictionary',*payload['tables']},'Workbook sheets differ from Source Data specification')
    for name,t in payload['tables'].items():
        rows=list(wb[name].iter_rows(values_only=True))
        check(list(rows[0])==t['columns'],'Workbook header changed: '+name)
        check(len(rows)-1==len(t['rows']),'Workbook row count changed: '+name)
        for col in t['columns']:check((name,col) in docs,'Undocumented Source Data column: '+name+'.'+col)
        # Roundtrip every statistical and raw value; blank is unavailable, never zero.
        for i,(observed,expected) in enumerate(zip(rows[1:],t['rows']),2):
            for a,b in zip(observed,expected):
                equal=(a==b) if not isinstance(b,(int,float)) or isinstance(b,bool) else isinstance(a,(int,float)) and np.isclose(a,b,rtol=1e-12,atol=1e-12)
                if not equal:errors.append(f'Workbook roundtrip mismatch: {name} row{i}');break
    wb.close()
    doc=PAPER/'Galaxy_agents_original_layout_manuscript.docx'
    with zipfile.ZipFile(doc) as z:
        for n in z.namelist():
            if n.endswith(('.xml','.rels')):ET.fromstring(z.read(n))
        root=ET.fromstring(z.read('word/document.xml'))
        text='\n'.join(n.text or '' for n in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        check(not any(s in text for s in ['{{','[@','[[FIGURES]]','[[REFERENCES]]']),'DOCX has unresolved build tokens')
        check(len([n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')])==6,'DOCX must embed six main figures')
        check('results pending' in text.lower(),'DOCX lacks intervention placeholder')
    text_filled=re.sub(r'\{\{([\w.-]+)\}\}',lambda m:numbers[m[1]],md)
    def count(s):return len([w for w in re.sub(r'\[@[^\]]+\]','',s).split() if re.search('[A-Za-z0-9]',w)])
    abstract=text_filled.split('## Abstract\n')[1].split('## [Introduction]')[0]
    maintext=text_filled.split('## [Introduction]\n')[1].split('## Figure legends')[0]
    abstract_words=count(abstract)
    main_words=count('\n'.join(l for l in maintext.splitlines() if not l.startswith('#')))
    check(abstract_words<=150 and main_words<=3000,'Analysis abstract/main word limits exceeded')
    sources=[NARR/'narrative_common.py',NARR/'build_docx.js',NARR/'references.json',NARR/'author_metadata.json',NARR/'requirements.txt',NARR/'package.json',
             ROOT/'BixBench50_CompBio_analysis/analysis.json',ROOT/'manuscript_material/source_data/derived/run_summaries.jsonl.gz',
             ROOT/'manuscript_material/on_demand/Source_Data_OD_Fig4.xlsx',ROOT/'manuscript_material/scripts/style.py',
             ROOT/'individual_error_analysis.md',ROOT/'IWC/iwc_scientific_audit.json']
    sources+=sorted((NARR/'derived').rglob('*.json'))+sorted((NARR/'derived').rglob('*.csv'))+sorted((NARR/'derived').rglob('*.gz'))
    ev=pd.read_csv(PAPER/'analysis/rigor_evidence.csv')
    for _,r in ev.iterrows():
        source=ROOT/r.source_file
        if source.exists():
            opener=gzip.open if source.suffix=='.gz' else open
            with opener(source,'rt') as f:
                line=next((l for i,l in enumerate(f,1) if i==r.decompressed_line),'')
            check(hashlib.sha256(line.encode()).hexdigest()==r.line_sha256,'Trace line hash changed: '+r.source_file)
            sources.append(source)
        else:check(False,'Missing trace: '+r.source_file)
    rigor=json.loads((PAPER/'analysis/rigor_results.json').read_text())
    sources.extend(ROOT/p for p in rigor['source_sha256'])
    versions={n:importlib.metadata.version(n) for n in ['numpy','pandas','matplotlib','scipy','openpyxl','pypdf']}
    versions['python']=sys.version.split()[0]
    versions['node']=subprocess.check_output([os.environ.get('NODE_BINARY','node'),'--version'],text=True).strip()
    versions['node_libraries']=json.loads(subprocess.check_output([os.environ.get('NODE_BINARY','node'),'-e',
        'const fs=require("fs"),p=require("path");function v(n){let d=p.dirname(require.resolve(n));while(d!==p.dirname(d)){let f=p.join(d,"package.json");if(fs.existsSync(f)){let j=JSON.parse(fs.readFileSync(f,"utf8"));if(j.name===n)return j.version;}d=p.dirname(d);}return "unavailable";}process.stdout.write(JSON.stringify({docx:v("docx"),artifact_tool:v("@oai/artifact-tool")}));'],text=True))
    generated=[p for p in PAPER.rglob('*') if p.is_file() and p.suffix in {'.md','.py','.mjs','.sh','.csv','.gz','.json','.pdf','.png','.docx','.xlsx'}
               and '__pycache__' not in p.parts and p.name not in {'release_manifest.json','package_validation.json','visual_validation.json'}]
    manifest={'schema_version':1,'versions':versions,'archived_inputs':{str(p.relative_to(ROOT)):sha(p) for p in sorted(set(sources)) if p.exists()},
              'package_files':{str(p.relative_to(ROOT)):sha(p) for p in sorted(generated)},
              'limits':'Hashes establish a local archive build, not scientific validity or independent replay. No private reference key accessed. Intervention result pending.'}
    report={'status':'failed' if errors else 'draft_checks_passed','errors':errors,'abstract_words':abstract_words,'main_words':main_words,
            'counts':{'endpoint_runs':len(runs),'replicate_sets':len(sets),'audited_cases':len(cases),'census_failing_binary_tasks':int(cases.primary_failing_task.sum()),
                      'figures':len(pdfs),'workbook_sheets':len(payload['tables'])+2},
            'pending_author_fields':sorted(set(pending)),'submission_ready':False,
            'pending_studies':['Independent official-key evaluation','Blinded expert audit','Artifact replay and second deployment','Token-reduction intervention'],
            'visual_QA':'Recorded separately after rendered-page inspection; this script does not establish layout quality.'}
    (PAPER/'release_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    (PAPER/'package_validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))
    return bool(errors)

if __name__=='__main__':raise SystemExit(main())
