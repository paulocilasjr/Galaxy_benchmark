"""Validate the local manuscript release, without running agents or contacting servers.

Default mode checks generated artifacts and reports unresolved submission fields.
--submission additionally fails on unresolved author fields and missing release attestations.
"""
import argparse
import csv
import gzip
import hashlib
import importlib.metadata
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from pypdf import PdfReader

NARR = Path(__file__).resolve().parents[1]
ROOT = NARR.parent
PAPERS = ('user-oriented', 'galaxy-oriented')


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--submission', action='store_true')
    args = parser.parse_args()
    errors, warnings, pending, counts = [], [], [], {}
    author_path = NARR / 'author_metadata.json'
    author = json.loads(author_path.read_text()) if author_path.exists() else {}
    fields = author.get('fields', {})
    refs = json.loads((NARR / 'references.json').read_text())
    for paper in PAPERS:
        folder = NARR / paper
        numbers = json.loads((folder / 'numbers.json').read_text())
        provenance = json.loads((folder / 'numbers_provenance.json').read_text())
        for template in [folder / 'manuscript.md', *sorted((folder / 'supplementary').glob('*.md'))]:
            text = template.read_text()
            for key in re.findall(r'\{\{([\w.-]+)\}\}', text):
                if key not in numbers or key not in provenance:
                    errors.append(f'{template.relative_to(ROOT)}: missing number/provenance {key}')
            for group in re.findall(r'\[@([^\]]+)\]', text):
                for key in group.split(';'):
                    if key.strip().lstrip('@') not in refs:
                        errors.append(f'{template.relative_to(ROOT)}: unknown citation {key}')
            for field in re.findall(r'\[Authors:\s*([^\]]+)\]', text):
                if not isinstance(fields.get(field.strip()), str) or not fields[field.strip()].strip():
                    pending.append({'file': str(template.relative_to(ROOT)), 'field': field.strip()})
        pdfs = sorted((folder / 'figures').glob('*.pdf'))
        main_figs = [p for p in pdfs if re.fullmatch(r'Fig[1-6]\.pdf', p.name)]
        ed_figs = [p for p in pdfs if p.name.startswith('ED_Fig')]
        if len(main_figs) != 6 or len(ed_figs) > 10:
            errors.append(f'{paper}: main/Extended Data figure count invalid')
        for pdf in pdfs:
            reader = PdfReader(pdf)
            width = float(reader.pages[0].mediabox.width) * 25.4 / 72
            if len(reader.pages) != 1 or width > 180.01:
                errors.append(f'{pdf.relative_to(ROOT)}: page count or width invalid ({width:.3f} mm)')
            source = folder / 'source_data' / f'Source_Data_{pdf.stem}.xlsx'
            if not source.exists():
                errors.append(f'{paper}: missing Source Data for {pdf.stem}')
                continue
            wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
            if not {'README', 'dictionary'} <= set(wb.sheetnames):
                errors.append(f'{source.relative_to(ROOT)}: missing documentation sheets')
            else:
                dictionary = list(wb['dictionary'].iter_rows(min_row=2, values_only=True))
                documented = {(row[0], row[1]) for row in dictionary if row[2]}
                if any(str(row[2]).startswith('Value shown in the figure') for row in dictionary):
                    errors.append(f'{source.relative_to(ROOT)}: generic column descriptions remain')
                for sheet in wb.sheetnames:
                    if sheet in ('README', 'dictionary'):
                        continue
                    columns = next(wb[sheet].iter_rows(values_only=True), ())
                    for column in columns:
                        if column is not None and (sheet, column) not in documented:
                            errors.append(f'{source.relative_to(ROOT)}: undocumented column {sheet}.{column}')
            wb.close()
        docxs = [*folder.glob('*.docx'), *(folder / 'supplementary').glob('*.docx')]
        if len(docxs) != 2:
            errors.append(f'{paper}: expected one manuscript and one Supplementary Note DOCX')
        for docx in docxs:
            with zipfile.ZipFile(docx) as archive:
                for name in archive.namelist():
                    if name.endswith('.xml') or name.endswith('.rels'):
                        ET.fromstring(archive.read(name))
                root = ET.fromstring(archive.read('word/document.xml'))
                text = '\n'.join(n.text or '' for n in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
                if '{{' in text or '[@' in text or '[[FIGURES]]' in text:
                    errors.append(f'{docx.relative_to(ROOT)}: unresolved build token')
        counts[paper] = {'main_figures': len(main_figs), 'extended_data_figures': len(ed_figs), 'docx_files': len(docxs),
                         'numeric_values': len(numbers)}
    with gzip.open(NARR / 'derived/galaxy_calls/calls.csv.gz', 'rt') as stream:
        calls = list(csv.DictReader(stream))
    galaxy = [r for r in calls if r['galaxy_server'] == 'True']
    failed = [r for r in galaxy if r['extract_fail'] == 'True' or r['call_status'] == 'failed']
    zero = [r for r in galaxy if r['tool'] == 'run_galaxy_tool_and_wait' and r['result_status'] == 'ok'
            and r['prov_status'] == 'no_explicit_non_dataset_parameters' and float(r['checked_parameter_count'] or -1) == 0]
    extra = [r for r in galaxy if r['extract_fail'] != 'True' and r['call_status'] == 'failed']
    counts['interface'] = {'all_calls': len(calls), 'galaxy_calls': len(galaxy), 'other_calls': len(calls)-len(galaxy),
                           'failed_galaxy_calls': len(failed), 'additional_exceptions': len(extra), 'ok_zero_checks': len(zero),
                           'parameter_mismatches': sum(r['prov_status'] == 'mismatch' and r['tool'] == 'run_galaxy_tool_and_wait' for r in galaxy),
                           'parameter_mismatches_all_operations': sum(r['prov_status'] == 'mismatch' for r in galaxy)}
    # Report the known independent audit invariants, rather than silently accepting changed populations.
    expected = {'galaxy_calls': 69505, 'failed_galaxy_calls': 7987, 'additional_exceptions': 604,
                'ok_zero_checks': 1140, 'parameter_mismatches': 4999}
    for key, value in expected.items():
        if counts['interface'][key] != value:
            errors.append(f'Independent archive invariant changed: {key}={counts["interface"][key]}, expected {value}')
    if len(pending):
        warnings.append(f'{len(pending)} author-supplied field occurrences remain unresolved; see author_metadata.json')
    for flag in ('human_audit_verified', 'release_identifiers_verified', 'compbio_key_access_terms_verified'):
        if not author.get('attestations', {}).get(flag, False):
            warnings.append(f'Submission attestation pending: {flag}')
    if args.submission and warnings:
        errors.extend(warnings)
    report = {'status': 'fail' if errors else 'pass_with_pending_submission_fields' if warnings else 'pass',
              'counts': counts, 'errors': errors, 'warnings': warnings, 'pending_author_fields': pending,
              'scope': 'Archived artifact regeneration and structural checks; no scientific reruns, human review or intervention validation.'}
    report['mode'] = 'submission' if args.submission else 'artifact'
    report_name = 'submission_validation.json' if args.submission else 'package_validation.json'
    (NARR / report_name).write_text(json.dumps(report, indent=2) + '\n')
    source_files = [ROOT / 'BixBench50_CompBio_analysis/analysis.json', ROOT / 'manuscript_material/source_data/figure_data.json',
                    ROOT / 'manuscript_material/source_data/derived/run_summaries.jsonl.gz',
                    ROOT / 'IWC/iwc_scientific_audit.json', *sorted((ROOT / 'manuscript_material/on_demand').glob('Source_Data_OD_Fig*.xlsx'))]
    source_files += sorted((NARR / 'derived').rglob('*.json')) + sorted((NARR / 'derived').rglob('*.csv')) + sorted((NARR / 'derived').rglob('*.gz'))
    source_files += sorted((ROOT / 'IWC/analysis').glob('*/source_snapshots/huggingface_traces/files/*/evaluation.json'))
    source_files += [ROOT / 'manuscript_material/scripts/style.py', ROOT / 'individual_error_analysis.md',
                     ROOT / 'analysis_reports/galaxy_improvement_20260924/v2_trace_friction/cat_tool.json',
                     ROOT / 'analysis_reports/galaxy_improvement_20260924/v2_trace_friction/ledger.json']
    generated = [p for p in NARR.rglob('*') if p.is_file() and p.suffix in ('.md', '.py', '.js', '.sh', '.json', '.txt', '.pdf', '.docx', '.xlsx')
                 and p.name not in ('release_manifest.json', 'package_validation.json', 'submission_validation.json')]
    versions = {name: importlib.metadata.version(name) for name in ('numpy', 'pandas', 'matplotlib', 'scipy', 'openpyxl', 'pypdf')}
    node_version = 'not verified'
    try:
        node_version = subprocess.check_output([os.environ.get('NODE_BINARY', 'node'), '-p',
            'JSON.parse(require("fs").readFileSync(require("path").join(require("path").dirname(require.resolve("docx")), "../package.json"), "utf8")).version'],
            text=True, stderr=subprocess.DEVNULL).strip()
    except (OSError, subprocess.CalledProcessError):
        warnings.append('Node document-library version unavailable to validator')
    manifest = {'schema_version': 1, 'python': sys.version.split()[0], 'packages': versions,
                'node_docx_version': node_version,
                'archived_inputs': {str(p.relative_to(ROOT)): digest(p) for p in source_files if p.exists()},
                'package_files': {str(p.relative_to(ROOT)): digest(p) for p in sorted(generated)},
                'private_compbio_key': 'Not accessed; frozen per-run grades reused. Independent regeneration requires authorized evaluator access.',
                'limits': 'Hashes establish frozen inputs and artifacts, not biological correctness or successful scientific replay.'}
    (NARR / 'release_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps({'status': report['status'], 'counts': counts, 'errors': errors, 'pending_fields': len(pending)}, indent=2))
    return bool(errors)


if __name__ == '__main__':
    raise SystemExit(main())
